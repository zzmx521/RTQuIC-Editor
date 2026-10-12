#!/usr/bin/env python3
"""RT-QuIC Data Editor — Prism-style with adjustable font sizes."""
import sys, os, json, csv, io, shutil, threading, subprocess, datetime
import numpy as np
from collections import defaultdict
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QSplitter, QTreeWidget, QTreeWidgetItem,
    QTabWidget, QTableWidget, QTableWidgetItem, QVBoxLayout, QHBoxLayout,
    QWidget, QPushButton, QLabel, QFileDialog, QMessageBox, QStatusBar,
    QFrame, QInputDialog, QAbstractItemView, QMenu, QDialog, QSlider, QDialogButtonBox,
    QPlainTextEdit, QCheckBox, QLineEdit, QGridLayout, QListWidget, QListWidgetItem, QScrollArea, QComboBox, QSizePolicy
)
from PyQt5.QtCore import Qt, pyqtSignal, QUrl, QTimer, QObject, pyqtSlot, QStandardPaths
from PyQt5.QtGui import QFont, QColor, QDesktopServices
from PyQt5.QtWebEngineWidgets import QWebEngineView, QWebEnginePage, QWebEngineProfile
import openpyxl
import icons

APP_DIR = os.path.dirname(os.path.abspath(__file__))


def _resource_path(name):
    """Return a bundled resource path both in source and PyInstaller builds."""
    root = getattr(sys, '_MEIPASS', APP_DIR)
    return os.path.join(root, name)


BUILDER = _resource_path('build_viewer.py')
PLOTLY_SRC = _resource_path('plotly.min.js')

# Keep the established macOS data location. On Windows, store user data in
# Documents so workspaces and generated views remain easy to find and back up.
if sys.platform.startswith('win'):
    _documents = QStandardPaths.writableLocation(QStandardPaths.DocumentsLocation)
    EDITOR_DIR = os.path.join(_documents or os.path.expanduser('~'), 'RT-QuIC-Editor')
else:
    EDITOR_DIR = os.path.expanduser('~/rt-viewer')

_desktop = QStandardPaths.writableLocation(QStandardPaths.DesktopLocation)
DESKTOP_DIR = _desktop or os.path.join(os.path.expanduser('~'), 'Desktop')
os.makedirs(EDITOR_DIR, exist_ok=True)
INIT_ROWS, INIT_COLS = 50, 30
SETTINGS_FILE = os.path.join(EDITOR_DIR, 'font_settings.json')

# When frozen (PyInstaller .app), prefer the copies bundled inside the bundle.
if getattr(sys, 'frozen', False):
    _meipass = getattr(sys, '_MEIPASS', os.path.dirname(sys.executable))
    for cand in (os.path.join(_meipass, 'build_viewer.py'), os.path.join(_meipass, 'plotly.min.js')):
        if os.path.exists(cand):
            if cand.endswith('build_viewer.py'):
                BUILDER = cand
            else:
                PLOTLY_SRC = cand

sys.path.insert(0, os.path.dirname(BUILDER))
import importlib.util
spec = importlib.util.spec_from_file_location('build_viewer', BUILDER)
bv = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bv)

GRAY = QColor('#e8eaef')
GRAY2 = QColor('#f0f1f5')


def _open_local_path(path):
    """Open a local file or directory with the operating system default app."""
    QDesktopServices.openUrl(QUrl.fromLocalFile(os.path.abspath(path)))


def _reveal_local_path(path):
    """Reveal a local file in Finder/Explorer, or open its containing folder."""
    path = os.path.abspath(path)
    try:
        if sys.platform.startswith('win'):
            if os.path.isdir(path):
                subprocess.Popen(['explorer.exe', path])
            else:
                subprocess.Popen(['explorer.exe', '/select,', os.path.normpath(path)])
        elif sys.platform == 'darwin':
            subprocess.run(['open', path] if os.path.isdir(path) else ['open', '-R', path], check=False)
        else:
            target = path if os.path.isdir(path) else os.path.dirname(path)
            subprocess.Popen(['xdg-open', target])
    except Exception:
        _open_local_path(path if os.path.isdir(path) else os.path.dirname(path))


def _to_hours(t):
    """Convert a time value (number or text like '1 h 30 min' / '0:30:00') to decimal hours."""
    if t is None or t == '':
        return ''
    if isinstance(t, (int, float)):
        return float(t)
    s = str(t).strip()
    try:
        return float(s)
    except ValueError:
        pass
    hours = 0.0
    found = False
    parts = s.split()
    i = 0
    while i < len(parts):
        try:
            num = float(parts[i])
        except (ValueError, TypeError):
            i += 1
            continue
        unit = parts[i + 1].lower() if i + 1 < len(parts) else ''
        if unit.startswith('h'):
            hours += num; found = True; i += 2
        elif unit.startswith('min'):
            hours += num / 60.0; found = True; i += 2
        elif unit.startswith('s'):
            hours += num / 3600.0; found = True; i += 2
        else:
            hours += num; found = True; i += 1
    if found:
        return round(hours, 4)
    # HH:MM:SS or MM:SS
    p = s.split(':')
    if len(p) == 3:
        try:
            return round(float(p[0]) + float(p[1]) / 60 + float(p[2]) / 3600, 4)
        except ValueError:
            pass
    elif len(p) == 2:
        try:
            return round(float(p[0]) / 60 + float(p[1]) / 3600, 4)
        except ValueError:
            pass
    return s


# ============================== FONT SETTINGS =================================
def load_settings():
    d = {'title': 24, 'sidebar': 16, 'toolbar': 14, 'table': 12, 'header': 12, 'gen_btn': 28,
         'lang': 'zh', 'win_geo': None}
    try:
        if os.path.exists(SETTINGS_FILE):
            d.update(json.load(open(SETTINGS_FILE)))
    except: pass
    return d

def save_settings(s):
    with open(SETTINGS_FILE, 'w') as f: json.dump(s, f)

# 备注模板：{模板名: [[字段名, 值], ...]}，应用级共享（不随工作区）
NOTE_TPL_FILE = os.path.join(EDITOR_DIR, 'note_templates.json')

def load_note_templates():
    try:
        with open(NOTE_TPL_FILE, encoding='utf-8') as f:
            d = json.load(f)
        return d if isinstance(d, dict) else {}
    except Exception:
        return {}

def save_note_templates(d):
    try:
        with open(NOTE_TPL_FILE, 'w', encoding='utf-8') as f:
            json.dump(d, f, ensure_ascii=False, indent=1)
    except Exception:
        pass

def T(zh, en):
    return zh if load_settings().get('lang', 'zh') == 'zh' else en

class Fonts:
    def __init__(self):
        self._s = load_settings()
    @staticmethod
    def _sys(size):
        # Apple 风格：使用 macOS 系统字体（SF Pro），不指定族名即取应用默认字体
        f = QFont(); f.setPointSize(size); return f
    def get(self, key): return self._sys(self._s.get(key, 12))
    def get_bold(self, key):
        f = self._sys(self._s.get(key, 12)); f.setBold(True); return f
    def size(self, key): return self._s.get(key, 12)
    def reload(self): self._s = load_settings()

FONTS = Fonts()

class SettingsDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle('Font Size Settings'); self.setMinimumWidth(400)
        layout = QVBoxLayout(self)
        self._sliders = {}
        labels = {'title': 'Title', 'sidebar': 'Sidebar', 'toolbar': 'Toolbar / Status',
                  'table': 'Data Table', 'header': 'Table Headers', 'gen_btn': '📈 Button'}
        for key, label in labels.items():
            layout.addWidget(QLabel(f'{label}:'))
            slider = QSlider(Qt.Horizontal); slider.setRange(8, 48)
            slider.setValue(FONTS._s.get(key, 12)); slider.setTickPosition(QSlider.TicksBelow)
            layout.addWidget(slider)
            val_lbl = QLabel(f'{slider.value()}pt'); layout.addWidget(val_lbl)
            slider.valueChanged.connect(lambda v, l=val_lbl: l.setText(f'{v}pt'))
            self._sliders[key] = slider
        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel | QDialogButtonBox.Reset)
        btns.accepted.connect(self._apply)
        btns.rejected.connect(self.reject)
        def reset():
            for k, s in self._sliders.items(): s.setValue({'title':24,'sidebar':16,'toolbar':14,'table':12,'header':12,'gen_btn':28}[k])
        btns.button(QDialogButtonBox.Reset).clicked.connect(reset)
        layout.addWidget(btns)

    def _apply(self):
        for k, s in self._sliders.items(): FONTS._s[k] = s.value()
        save_settings(FONTS._s); self.accept()


# ============================== DATA MODEL ====================================
class SheetModel:
    def __init__(self, rows=INIT_ROWS, cols=INIT_COLS):
        self.sample_names = [''] * rows; self.time_points = [''] * cols
        self.data = [[None] * cols for _ in range(rows)]
        self.excluded = []
    def ensure(self, r, c):
        while len(self.sample_names) < r: self.sample_names.append('')
        while len(self.time_points) < c: self.time_points.append('')
        while len(self.data) < r: self.data.append([None] * max(c, len(self.time_points)))
        for row in self.data:
            while len(row) < max(c, len(self.time_points)): row.append(None)
    @property
    def rows(self): return len(self.sample_names)
    @property
    def cols(self): return len(self.time_points)
    def get(self, r, c):
        if 0 <= r < len(self.data) and 0 <= c < len(self.data[r]):
            v = self.data[r][c]; return v if v is not None else None
        return None
    def set(self, r, c, v): self.ensure(r + 1, c + 1); self.data[r][c] = v
    def to_csv(self):
        lines = [['Time'] + [str(t) for t in self.time_points]]
        for i, name in enumerate(self.sample_names):
            row = [name] + [str(self.data[i][j]) if self.get(i, j) is not None else '' for j in range(len(self.time_points))]
            lines.append(row)
        return '\n'.join(','.join(self._q(c) for c in r) for r in lines)
    def from_text(self, text):
        lines = [l.rstrip('\r') for l in text.split('\n') if l.strip()]
        if len(lines) < 2: return
        delim = '\t' if '\t' in lines[0] else ','
        matrix = [self._split(l, delim) for l in lines]
        self.time_points = []
        for v in matrix[0][1:]:
            v = v.strip()
            try: self.time_points.append(float(v))
            except ValueError: self.time_points.append(v if v else '')
        self.sample_names, self.data = [], []
        for row in matrix[1:]:
            self.sample_names.append(row[0].strip() if row else '')
            rdata = []
            for v in row[1:]:
                v = v.strip()
                if not v: rdata.append(None)
                else:
                    try: rdata.append(float(v))
                    except ValueError: rdata.append(None)
            while len(rdata) < len(self.time_points): rdata.append(None)
            self.data.append(rdata)
    def from_json(self, text):
        d = json.loads(text)
        self.time_points = d.get('timePoints', [''] * INIT_COLS)
        self.sample_names = d.get('sampleNames', [''] * INIT_ROWS)
        self.data = d.get('data', [[None] * INIT_COLS for _ in range(INIT_ROWS)])
    def to_json(self):
        return json.dumps({'timePoints': self.time_points, 'sampleNames': self.sample_names, 'data': self.data})
    def _split(self, line, delim):
        if delim == '\t': return line.split('\t')
        result, cur, q = [], '', False
        for ch in line:
            if q:
                if ch == '"': q = False
                else: cur += ch
            else:
                if ch == '"': q = True
                elif ch == ',': result.append(cur); cur = ''
                else: cur += ch
        result.append(cur); return result
    def _q(self, s):
        if any(c in s for c in [',', '"', '\n']): return '"' + s.replace('"', '""') + '"'
        return s


# ============================== TREE WIDGET ===================================
class TreeWidget(QTreeWidget):
    """QTreeWidget with a drop callback for drag-reorder persistence."""
    def __init__(self, on_drop=None, parent=None):
        super().__init__(parent); self._on_drop = on_drop
    def dropEvent(self, event):
        super().dropEvent(event)
        if self._on_drop: self._on_drop()


# ============================== SHEET TABLE WIDGET ============================
class SheetWidget(QWidget):
    edit_done = pyqtSignal(str)
    dirty_changed = pyqtSignal(str)
    def __init__(self, model, name, parent=None):
        super().__init__(parent); self.model = model; self._name = name; self._build()
        self._undo = []; self._redo = []; self._dirty = False
        self._last = self._snapshot()
        self._debounce = QTimer(self); self._debounce.setSingleShot(True)
        self._debounce.timeout.connect(lambda: self.edit_done.emit(self._name))
        self._table.itemChanged.connect(self._on_item_changed)
    def _build(self):
        layout = QVBoxLayout(self); layout.setContentsMargins(0, 0, 0, 0)
        class EnterTable(QTableWidget):
            def keyPressEvent(_, event):
                if event.key() in (Qt.Key_Return, Qt.Key_Enter):
                    cur = _.currentIndex()
                    if cur.isValid():
                        _.commitData(_.cellWidget(cur.row(), cur.column()))
                        nr = min(cur.row() + 1, _.rowCount() - 1)
                        _.setCurrentCell(nr, cur.column()); _.edit(_.currentIndex())
                    return
                super(EnterTable, _).keyPressEvent(event)
        self._table = EnterTable()
        self._table.setStyleSheet('QTableWidget{background:#fff;gridline-color:#e5e5ea;border:none;}QTableWidget::item:selected{background:#d6e9ff;color:#1d1d1f;}')
        self._table.setSelectionMode(QAbstractItemView.SingleSelection)
        self._table.setEditTriggers(QAbstractItemView.CurrentChanged)
        self._table.setTabKeyNavigation(True)
        self._table.setContextMenuPolicy(Qt.CustomContextMenu)
        self._table.customContextMenuRequested.connect(self._on_table_menu)
        self._table.horizontalHeader().hide(); self._table.verticalHeader().hide()
        self._scale_check = QCheckBox('×10³ 显示'); self._scale_check.setFont(FONTS.get('toolbar'))
        self._scale_check.toggled.connect(lambda _: self.sync())
        layout.addWidget(self._scale_check)
        layout.addWidget(self._table)
    def sync(self):
        nr, nc = self.model.rows, self.model.cols
        self._table.setRowCount(nr + 1); self._table.setColumnCount(nc + 1)
        self._table.blockSignals(True)
        tf = FONTS.get('table'); hf = FONTS.get('header')
        c = QTableWidgetItem(''); c.setBackground(GRAY); c.setFlags(Qt.NoItemFlags)
        self._table.setItem(0, 0, c)
        for j in range(nc):
            v = self.model.time_points[j]; txt = str(v) if (v is not None and v != '') else ''
            it = QTableWidgetItem(txt); it.setBackground(GRAY); it.setFont(hf); it.setTextAlignment(Qt.AlignCenter)
            self._table.setItem(0, j + 1, it)
        for i in range(nr):
            name = self.model.sample_names[i] or ''
            si = QTableWidgetItem(name); si.setBackground(GRAY); si.setFont(hf)
            self._table.setItem(i + 1, 0, si)
            _sc = 1000.0 if self._scale_check.isChecked() else 1.0
            for j in range(nc):
                v = self.model.get(i, j)
                txt = self._fmt(v / _sc) if v is not None else ''
                di = QTableWidgetItem(txt); di.setFont(tf); di.setTextAlignment(Qt.AlignCenter)
                self._table.setItem(i + 1, j + 1, di)
        for i in (self.model.excluded or []):
            if 0 <= i < nr:
                for j in range(nc + 1):
                    it = self._table.item(i + 1, j)
                    if it:
                        f = it.font(); f.setItalic(True); it.setFont(f)
                        it.setForeground(QColor('#999999'))
        self._table.setColumnWidth(0, 150)
        for j in range(nc): self._table.setColumnWidth(j + 1, 94)
        self._table.setRowHeight(0, FONTS.size('header') + 16)
        for i in range(nr): self._table.setRowHeight(i + 1, FONTS.size('table') + 16)
        self._table.blockSignals(False)
    def model_from_table(self):
        nr, nc = self.model.rows, self.model.cols
        for j in range(nc):
            it = self._table.item(0, j + 1)
            if it:
                txt = it.text().strip()
                try: self.model.time_points[j] = float(txt)
                except ValueError: self.model.time_points[j] = txt if txt else ''
        for i in range(nr):
            it = self._table.item(i + 1, 0)
            if it: self.model.sample_names[i] = it.text().strip()
            _sc = 1000.0 if self._scale_check.isChecked() else 1.0
            for j in range(nc):
                it = self._table.item(i + 1, j + 1)
                if it:
                    txt = it.text().strip()
                    if not txt: self.model.set(i, j, None)
                    else:
                        try: self.model.set(i, j, float(txt) * _sc)
                        except ValueError: pass
    def _fmt(self, v):
        if v is None: return ''
        if isinstance(v, float) and v == int(v): return str(int(v))
        if isinstance(v, float): s = f'{v:.6g}'; return s if '.' in s else s + '.0'
        return str(v)
    def _on_table_menu(self, pos):
        row = self._table.rowAt(pos.y())
        if row <= 0:
            return
        i = row - 1
        if i >= self.model.rows:
            return
        excluded = i in (self.model.excluded or [])
        menu = QMenu(self)
        act = menu.addAction('恢复此复孔（计入均值/SD）' if excluded else '剔除此复孔（不计入均值/SD）')
        if menu.exec_(self._table.mapToGlobal(pos)) == act:
            self.push_undo()
            if not self.model.excluded:
                self.model.excluded = []
            if excluded:
                self.model.excluded.remove(i)
            else:
                self.model.excluded.append(i)
            self.sync()
            self.commit_change()

    def _snapshot(self):
        return (list(self.model.time_points), list(self.model.sample_names),
                [list(r) for r in self.model.data])
    def _restore(self, snap):
        self.model.time_points = list(snap[0]); self.model.sample_names = list(snap[1])
        self.model.data = [list(r) for r in snap[2]]
        self.sync()
    def _on_item_changed(self, item):
        self.push_undo()
        self.commit_change()

    def push_undo(self):
        """Capture current state so the next change can be undone."""
        self.model_from_table()
        if self._last is not None:
            self._undo.append(self._last)
            if len(self._undo) > 60: self._undo.pop(0)
        self._redo.clear()

    def commit_change(self):
        """Record post-change state, mark dirty, schedule view refresh."""
        self.model_from_table()
        self._last = self._snapshot()
        self._dirty = True
        self.dirty_changed.emit(self._name)
        self._debounce.start(1500)
    def undo(self):
        if not self._undo: return
        self._redo.append(self._last)
        snap = self._undo.pop(); self._last = snap
        self._restore(snap); self._dirty = True
        self.dirty_changed.emit(self._name)
        self._debounce.start(1500)
    def redo(self):
        if not self._redo: return
        self._undo.append(self._last)
        snap = self._redo.pop(); self._last = snap
        self._restore(snap); self._dirty = True
        self.dirty_changed.emit(self._name)
        self._debounce.start(1500)
    def mark_clean(self):
        self._dirty = False
        self.dirty_changed.emit(self._name)
    def is_dirty(self):
        return self._dirty
    def table(self): return self._table


# ============================== VIEW WIDGET ===================================
class _ExternalPage(QWebEnginePage):
    """Open new-window requests (window.open / target=_blank) in the system browser."""
    def createWindow(self, _type):
        page = QWebEnginePage(self)
        def _open_external(url):
            s = url.toString()
            if s and not s.startswith('about:') and not s.startswith('data:') and not s.startswith('blob:'):
                QDesktopServices.openUrl(url)
        page.urlChanged.connect(_open_external)
        return page


class ViewWidget(QWidget):
    def __init__(self, html_path, parent=None, bare=False, view_name='', window=None):
        super().__init__(parent); self.html_path = html_path; self._bare = bare
        self._view_name = view_name; self._window = window
        self._build()
    def _build(self):
        layout = QVBoxLayout(self); layout.setContentsMargins(0, 0, 0, 0)
        self._web = QWebEngineView()
        self._web.setPage(_ExternalPage(self._web))
        self._web.setUrl(self._url()); layout.addWidget(self._web)

    def _url(self):
        u = QUrl.fromLocalFile(self.html_path)
        if self._bare:
            u.setQuery('bare=1')
        return u

    def refresh_fonts(self):
        for btn in self.findChildren(QPushButton):
            btn.setFont(FONTS.get('toolbar'))

    def reload(self):
        self._web.setUrl(self._url())


# ============================== MAIN WINDOW ===================================
class ZoomableView(QWidget):
    """A view cell with a stepless (continuous) zoom bar for multiview display.

    备注编辑统一由应用工具栏"📝 备注"按钮承担（视图页签下自动解析关联数据表），
    此处不放置备注按钮，避免重复入口。
    """
    def __init__(self, html_path, parent=None, bare=True, view_name='', window=None):
        super().__init__(parent)
        lay = QVBoxLayout(self); lay.setContentsMargins(0, 0, 0, 0); lay.setSpacing(2)
        bar = QHBoxLayout(); bar.setContentsMargins(4, 2, 4, 2); bar.setSpacing(4)
        minus = QPushButton('−'); minus.setFixedSize(22, 20); minus.setToolTip(T('缩小', 'Zoom out'))
        self._slider = QSlider(Qt.Horizontal); self._slider.setRange(50, 300); self._slider.setValue(100)
        plus = QPushButton('+'); plus.setFixedSize(22, 20); plus.setToolTip(T('放大', 'Zoom in'))
        self._lbl = QLabel('100%'); self._lbl.setMinimumWidth(42)
        self._lbl.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        bar.addWidget(minus); bar.addWidget(self._slider, 1); bar.addWidget(plus); bar.addWidget(self._lbl)
        lay.addLayout(bar)
        self._view = ViewWidget(html_path, bare=bare, view_name=view_name, window=window)
        lay.addWidget(self._view, 1)
        def _apply(v):
            z = v / 100.0
            self._view._web.setZoomFactor(z)
            self._lbl.setText(f'{v}%')
        self._slider.valueChanged.connect(_apply)
        minus.clicked.connect(lambda: self._slider.setValue(max(50, self._slider.value() - 10)))
        plus.clicked.connect(lambda: self._slider.setValue(min(300, self._slider.value() + 10)))

    def reload(self):
        self._view.reload()

    def web(self):
        return self._view._web


class MainWindow(QMainWindow):
    gen_done_signal = pyqtSignal(str, str)
    auto_gen_done_signal = pyqtSignal(str, str, str)  # path, name, sheet

    def __init__(self):
        super().__init__()
        self._sheets = {}; self._sheet_widgets = {}; self._views = {}; self._view_widgets = {}; self._view_sheet = {}; self._notes = {}
        self._building = False; self._sheet_counter = 0
        self._last_metrics_threshold = 1000
        self._multiview_widgets = []; self._multiview_tab = None
        self.gen_done_signal.connect(self._on_gen_done)
        self.auto_gen_done_signal.connect(self._on_auto_gen_done)
        try:
            QWebEngineProfile.defaultProfile().downloadRequested.connect(self._save_download)
        except Exception:
            pass
        self._init_ui(); self._load_workspace()
        self._ensure_plotly(); self._setup_shortcuts()
        geo = load_settings().get('win_geo')
        if geo and len(geo) == 4:
            try: self.setGeometry(*geo)
            except Exception: pass
        self._apply_lang()

    def _init_ui(self):
        self._tb_buttons = []
        self.setWindowTitle(T('RT-QuIC 数据编辑器', 'RT-QuIC Data Editor')); self.setGeometry(80, 60, 1400, 850); self.setMinimumSize(1000, 600)
        self.setStyleSheet('QMainWindow{background:#ececee;}')
        splitter = QSplitter(Qt.Horizontal); self.setCentralWidget(splitter)

        # Sidebar
        sidebar = QWidget(); sidebar.setMinimumWidth(160); sidebar.setMaximumWidth(400)
        sidebar.setStyleSheet('QWidget{background:#f5f5f7;border-right:1px solid #d2d2d7;}')
        sl = QVBoxLayout(sidebar); sl.setContentsMargins(6, 8, 6, 8); sl.setSpacing(4)
        title = QLabel(T('导航器', 'Navigator')); self._title_lbl = title; title.setFont(FONTS.get('title'))
        title.setStyleSheet('color:#007AFF;border:none;padding:4px;'); sl.addWidget(title)

        self._tree = TreeWidget(self._on_tree_drop); self._tree.setHeaderHidden(True); self._tree.setFont(FONTS.get('sidebar'))
        self._tree.setAnimated(True)
        self._tree.setStyleSheet('QTreeWidget{border:none;background:transparent;}QTreeWidget::item{padding:4px 2px;border-radius:5px;}QTreeWidget::item:selected{background:#007AFF;color:#fff;}QTreeWidget::item:hover{background:rgba(0,0,0,0.045);}')
        self._tree.itemClicked.connect(self._on_tree_clicked)
        self._tree.setContextMenuPolicy(Qt.CustomContextMenu)
        self._tree.customContextMenuRequested.connect(self._on_tree_context_menu)
        self._tree.setDragEnabled(True); self._tree.setAcceptDrops(True)
        self._tree.setDropIndicatorShown(True)
        self._tree.setDragDropMode(QAbstractItemView.InternalMove)
        sl.addWidget(self._tree)

        self._data_root = QTreeWidgetItem(['  ' + T('数据', 'Data')]); self._data_root.setIcon(0, icons.make_icon('table', '#6e6e73', 16)); self._data_root.setFont(0, FONTS.get('sidebar'))
        self._tree.addTopLevelItem(self._data_root)
        self._new_item = QTreeWidgetItem(['  ' + T('新建表', 'New Sheet')]); self._new_item.setIcon(0, icons.make_icon('plus', '#6e6e73', 14)); self._new_item.setData(0, Qt.UserRole, '__new__'); self._new_item.setFont(0, FONTS.get('sidebar'))
        self._new_item.setFlags(self._new_item.flags() & ~Qt.ItemIsDragEnabled)
        self._data_root.addChild(self._new_item); self._data_root.setExpanded(True)

        self._view_root = QTreeWidgetItem(['  ' + T('视图', 'View')]); self._view_root.setIcon(0, icons.make_icon('eye', '#6e6e73', 16)); self._view_root.setFont(0, FONTS.get('sidebar'))
        self._tree.addTopLevelItem(self._view_root);
        self._load_item = QTreeWidgetItem(['  ' + T('加载视图', 'Load View')]); self._load_item.setIcon(0, icons.make_icon('doc_plus', '#6e6e73', 14)); self._load_item.setData(0, Qt.UserRole, '__load_view__'); self._load_item.setFont(0, FONTS.get('sidebar'))
        self._load_item.setFlags(self._load_item.flags() & ~Qt.ItemIsDragEnabled)
        self._view_root.addChild(self._load_item); self._view_root.setExpanded(True)
        self._tree.currentItemChanged.connect(self._sync_selection_bold)
        splitter.addWidget(sidebar)

        # Right side
        right = QWidget(); rl = QVBoxLayout(right); rl.setContentsMargins(0, 0, 0, 0); rl.setSpacing(0)
        tb = QFrame(); tb.setFixedHeight(46)
        tb.setStyleSheet('QFrame{background:#f6f6f6;border-bottom:1px solid #d2d2d7;}')
        tbl = QHBoxLayout(tb); tbl.setContentsMargins(8, 4, 8, 4); tbl.setSpacing(4)
        self._tb_btn('导入', self._import, tbl, icon='folder', en='Import')
        self._tb_btn('导出CSV', self._export_csv, tbl, icon='export', en='Export CSV')
        self._tb_btn('导出JSON', self._export_json, tbl, icon='braces', en='Export JSON')
        self._tb_sep(tbl)
        self._tb_btn('+5行', lambda: self._add(5, 0), tbl, icon='rows', en='+5 Rows')
        self._tb_btn('+5列', lambda: self._add(0, 5), tbl, icon='cols', en='+5 Cols')
        self._tb_btn('清空', self._clear, tbl, icon='trash', en='Clear')
        self._tb_btn('历史', self._show_history, tbl, icon='clock', en='History')
        self._tb_btn('撤销', self._undo, tbl, icon='undo', en='Undo')
        self._tb_btn('重做', self._redo, tbl, icon='redo', en='Redo')
        self._tb_sep(tbl)
        self._tb_btn('导出指标', self._export_metrics, tbl, icon='chartbar', en='Metrics')
        self._tb_btn('备注', self._show_notes, tbl, icon='note', en='Notes')
        self._tb_btn('报告', self._export_report, tbl, icon='doc', en='Report')
        self._tb_sep(tbl)
        self._more_btn = QPushButton(T('☰ 更多', '☰ More')); self._more_btn.setFont(FONTS.get('toolbar')); self._more_btn.setCursor(Qt.PointingHandCursor)
        self._more_btn.setStyleSheet('QPushButton{background:transparent;border:none;color:#1d1d1f;padding:5px 8px;border-radius:6px;}QPushButton:hover{background:#e8e8ed;}QPushButton:pressed{background:#dcdce1;}')
        mm = QMenu(self._more_btn)
        mm.addAction(T('批量导入...', 'Batch Import...'), self._import_batch)
        mm.addAction(T('批量报告', 'Batch Report'), self._export_batch_report)
        mm.addAction(T('多视图同屏...', 'Multi-View Compare...'), self._show_multiview)
        mm.addAction(T('模板...', 'Templates...'), self._show_templates)
        mm.addAction(T('填充时间序列...', 'Fill Time Series...'), self._fill_time_series)
        mm.addSeparator()
        mm.addAction(T('English / 中文', 'English / 中文'), self._toggle_lang)
        self._more_btn.setMenu(mm)
        tbl.addWidget(self._more_btn)
        self._tb_sep(tbl)
        self._tb_btn('设置', self._show_settings, tbl, icon='gear', en='Settings')
        self._tb_sep(tbl)
        self._gen_btn = QPushButton(); self._gen_btn.setFont(FONTS.get('gen_btn'))
        self._gen_btn.setIcon(icons.make_icon('sparkles', '#ffffff', 22))
        self._gen_btn.setFixedSize(52, 38); self._gen_btn.setCursor(Qt.PointingHandCursor)
        self._gen_btn.setStyleSheet('QPushButton{background:#007AFF;border:none;border-radius:19px;}QPushButton:hover{background:#0071e3;}QPushButton:pressed{background:#0062cc;}')
        self._gen_btn.clicked.connect(self._generate); self._gen_btn.setToolTip('Generate Visualization')
        tbl.addWidget(self._gen_btn); tbl.addStretch()
        rl.addWidget(tb)

        self._tabs = QTabWidget(); self._tabs.setFont(FONTS.get('toolbar'))
        self._tabs.setStyleSheet('QTabWidget::pane{border:none;background:#f5f5f7;}QTabBar::tab{padding:6px 16px;font-weight:bold;}QTabBar::tab:selected{background:#f5f5f7;border-bottom:2px solid #007AFF;}')
        self._tabs.setTabsClosable(True); self._tabs.tabCloseRequested.connect(self._close_tab)
        self._tabs.tabBar().hide()
        rl.addWidget(self._tabs)
        splitter.addWidget(right); splitter.setSizes([220, 1180])

        self._status = QStatusBar(); self._status.setFont(FONTS.get('toolbar'))
        self._status.setStyleSheet('QStatusBar{background:#f6f6f6;border-top:1px solid #d2d2d7;color:#6e6e73;}')
        self._status.showMessage('Ready'); self.setStatusBar(self._status)
        self._status_info = QLabel('当前: —'); self._status_info.setFont(FONTS.get('toolbar'))
        self._status.addPermanentWidget(self._status_info)

    def _refresh_ui_fonts(self):
        """Re-apply fonts after settings change."""
        tf = FONTS.get('toolbar')
        self._gen_btn.setFont(FONTS.get('gen_btn'))
        for w in self.findChildren(QPushButton):
            if w is not self._gen_btn: w.setFont(FONTS.get('toolbar'))
        self._tabs.setFont(tf); self._status.setFont(tf); self._tree.setFont(FONTS.get('sidebar'))
        self._data_root.setFont(0, FONTS.get('sidebar')); self._view_root.setFont(0, FONTS.get('sidebar'))
        for i in range(self._data_root.childCount()):
            self._data_root.child(i).setFont(0, FONTS.get('sidebar'))
        for i in range(self._view_root.childCount()):
            self._view_root.child(i).setFont(0, FONTS.get('sidebar'))
        for sw in self._sheet_widgets.values(): sw.sync()
        for vw in self._view_widgets.values(): vw.refresh_fonts()

    def _show_settings(self):
        dlg = SettingsDialog(self)
        if dlg.exec_() == QDialog.Accepted:
            FONTS.reload(); self._refresh_ui_fonts()
            self._status.showMessage(T('字体设置已应用', 'Font settings applied'))

    def _apply_lang(self):
        for b, zh, en in getattr(self, '_tb_buttons', []):
            b.setText(T(zh, en))
        self.setWindowTitle(T('RT-QuIC 数据编辑器', 'RT-QuIC Data Editor'))
        if hasattr(self, '_title_lbl'):
            self._title_lbl.setText(T('导航器', 'Navigator'))
        if hasattr(self, '_data_root'):
            self._data_root.setText(0, '  ' + T('数据', 'Data'))
            self._view_root.setText(0, '  ' + T('视图', 'View'))
            self._new_item.setText(0, '  ' + T('新建表', 'New Sheet'))
            self._load_item.setText(0, '  ' + T('加载视图', 'Load View'))
        if hasattr(self, '_more_btn'):
            self._more_btn.setText(T('☰ 更多', '☰ More'))

    def _toggle_lang(self):
        s = load_settings()
        s['lang'] = 'en' if s.get('lang', 'zh') == 'zh' else 'zh'
        save_settings(s)
        self._apply_lang()
        self._status.showMessage(T('已切换到英文', 'Switched to Chinese') if s['lang'] == 'en' else '已切换到中文')

    def _show_templates(self):
        n = self._current_sheet_name()
        if not n: return
        items = [T('默认 50行×30列', 'Default 50x30'),
                 T('RT-QuIC 96孔 8行×12列', '96-well 8x12'),
                 T('滴定 3复孔 30行×15列', 'Titration 30x15'),
                 T('空白 20行×20列', 'Blank 20x20')]
        item, ok = QInputDialog.getItem(self, T('模板', 'Template'), T('选择模板:', 'Choose template:'), items, 0, False)
        if not ok: return
        rows, cols = {'0': (50, 30), '1': (8, 12), '2': (30, 15), '3': (20, 20)}[str(items.index(item))]
        sw = self._sheet_widgets[n]; sw.push_undo()
        m = self._sheets[n]
        m.ensure(rows, cols); sw.sync()
        sw.commit_change()
        self._status.showMessage(f'✅ 模板应用: {rows}×{cols}')

    def _fill_time_series(self):
        n = self._current_sheet_name()
        if not n: return
        start, ok1 = QInputDialog.getDouble(self, T('填充时间序列', 'Fill Time Series'), T('起始时间 (h):', 'Start (h):'), 0.0, 0, 1e9, 4)
        if not ok1: return
        step, ok2 = QInputDialog.getDouble(self, T('填充时间序列', 'Fill Time Series'), T('步长 (h):', 'Step (h):'), 0.5, 0.001, 1e9, 4)
        if not ok2: return
        count, ok3 = QInputDialog.getInt(self, T('填充时间序列', 'Fill Time Series'), T('点数:', 'Count:'), 20, 1, 1000)
        if not ok3: return
        sw = self._sheet_widgets[n]
        sw.push_undo()
        m = self._sheets[n]
        for j in range(count):
            m.time_points[j] = round(start + step * j, 4)
        sw.sync()
        sw.commit_change()
        self._status.showMessage(f'✅ 时间序列: {count} 点 ({start}–{round(start+step*(count-1),4)} h)')

    def _dirty_sheet_mark(self, name):
        pass

    def _grab_view_image(self, name):
        """Render a view to a base64 PNG (uses the open widget if present, else a temporary one)."""
        vw = self._view_widgets.get(name)
        temp = None
        if not vw:
            path = self._views.get(name)
            if not path or not os.path.exists(path):
                return None
            from PyQt5.QtCore import QEventLoop, QTimer
            vw = ViewWidget(path, view_name=name, window=self); temp = vw
            loop = QEventLoop()
            vw._web.loadFinished.connect(lambda ok: loop.quit())
            QTimer.singleShot(8000, loop.quit)
            loop.exec_()
        pm = vw._web.grab()
        if temp:
            temp.deleteLater()
        if pm.isNull():
            return None
        import base64 as _b64
        buf = __import__('io').BytesIO()
        pm.save(buf, 'PNG')
        return _b64.b64encode(buf.getvalue()).decode()

    def _show_multiview(self):
        try:
            self._show_multiview_impl()
        except Exception as e:
            import traceback; traceback.print_exc()
            QMessageBox.critical(self, T('多视图同屏', 'Multi-View'), T('打开多视图失败：', 'Failed to open multi-view: ') + str(e))

    def _show_multiview_impl(self):
        names = [vn for vn in self._views if os.path.exists(self._views.get(vn, ''))]
        if not names:
            QMessageBox.information(self, T('多视图同屏', 'Multi-View'), T('没有可用的视图。', 'No views available.')); return
        dlg = QDialog(self); dlg.setWindowTitle(T('选择要同屏显示的视图', 'Select views to display together')); dlg.setMinimumSize(380, 440)
        lay = QVBoxLayout(dlg)
        lst = QListWidget()
        for vn in names:
            it = QListWidgetItem(vn); it.setFlags(it.flags() | Qt.ItemIsUserCheckable)
            it.setCheckState(Qt.Checked if len(names) <= 6 else Qt.Unchecked)
            lst.addItem(it)
        lay.addWidget(lst)
        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.accepted.connect(dlg.accept); btns.rejected.connect(dlg.reject)
        lay.addWidget(btns)
        if dlg.exec_() != QDialog.Accepted:
            return
        chosen = [lst.item(i).text() for i in range(lst.count()) if lst.item(i).checkState() == Qt.Checked]
        if not chosen:
            QMessageBox.information(self, T('多视图同屏', 'Multi-View'), T('未选择视图。', 'No view selected.')); return
        self._open_multiview_tab(chosen)

    def _open_multiview_tab(self, chosen):
        try:
            self._open_multiview_tab_impl(chosen)
        except Exception as e:
            import traceback; traceback.print_exc()
            QMessageBox.critical(self, T('多视图同屏', 'Multi-View'), T('创建多视图页失败：', 'Failed to create multi-view: ') + str(e))

    def _open_multiview_tab_impl(self, chosen):
        for vn in chosen:
            self._ensure_view_fresh(vn)
        n = len(chosen)
        self._multiview_names = list(chosen)
        self._multiview_cols = 1 if n == 1 else 2
        wrap = QWidget()
        vl = QVBoxLayout(wrap); vl.setContentsMargins(6, 6, 6, 6)
        bar = QHBoxLayout()
        ref = QPushButton('🔄 ' + T('刷新', 'Refresh')); ref.clicked.connect(self._reload_multiview)
        rs = QPushButton(T('重新选择', 'Reselect')); rs.clicked.connect(self._show_multiview)
        cl = QPushButton('✖ ' + T('关闭', 'Close')); cl.clicked.connect(self._close_multiview)
        bar.addWidget(ref); bar.addWidget(rs); bar.addWidget(cl)
        bar.addWidget(QLabel(T('列数:', 'Columns:')))
        self._mv_cols_combo = QComboBox()
        for c in (1, 2, 3):
            self._mv_cols_combo.addItem(f'{c} {T("列", "col")}', c)
        self._mv_cols_combo.setCurrentIndex(self._multiview_cols - 1)
        self._mv_cols_combo.currentIndexChanged.connect(self._mv_cols_changed)
        bar.addWidget(self._mv_cols_combo); bar.addStretch()
        vl.addLayout(bar)
        self._multiview_grid_wrap = QWidget()
        self._multiview_grid = QGridLayout(self._multiview_grid_wrap)
        self._multiview_grid.setContentsMargins(0, 0, 0, 0)
        vl.addWidget(self._multiview_grid_wrap, 1)
        scroll = QScrollArea(); scroll.setWidgetResizable(True); scroll.setWidget(wrap)
        self._multiview_tab = scroll
        idx = self._tabs.addTab(scroll, '📊 ' + T('多视图', 'Multi-View'))
        self._tabs.setCurrentIndex(idx)
        self._multiview_widgets = []
        self._rebuild_multiview_grid(self._multiview_cols)

    def _mv_cols_changed(self, idx):
        cols = self._mv_cols_combo.itemData(idx) or 2
        self._multiview_cols = cols
        self._rebuild_multiview_grid(cols)

    def _rebuild_multiview_grid(self, cols):
        """Rearrange views into `cols` columns; keep a generous min width so curves don't compress."""
        while self._multiview_grid.count():
            it = self._multiview_grid.takeAt(0)
            if it.widget():
                self._multiview_grid.removeWidget(it.widget())
        if not self._multiview_widgets:
            for vn in self._multiview_names:
                self._multiview_widgets.append(ZoomableView(self._views[vn], bare=True, view_name=vn, window=self))
        rows_n = (len(self._multiview_widgets) + cols - 1) // cols
        for k, vw in enumerate(self._multiview_widgets):
            vw.setMinimumSize(320, 240)
            vw.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
            self._multiview_grid.addWidget(vw, k // cols, k % cols)
        # uniform equal distribution: every row and column gets the same stretch
        for c in range(cols):
            self._multiview_grid.setColumnStretch(c, 1)
        for r in range(rows_n):
            self._multiview_grid.setRowStretch(r, 1)

    def _reload_multiview(self):
        for k, vw in enumerate(self._multiview_widgets):
            vw.reload()
        self._status.showMessage('🔄 多视图已刷新')

    def _close_multiview(self):
        if self._multiview_tab is None:
            return
        idx = self._tabs.indexOf(self._multiview_tab)
        if idx >= 0:
            self._tabs.removeTab(idx)
        self._multiview_widgets = []; self._multiview_tab = None

    def _export_batch_report(self):
        sheets = [n for n in self._sheets if n in self._sheet_widgets]
        if not sheets:
            QMessageBox.information(self, T('批量报告', 'Batch Report'), T('没有数据表。', 'No sheets.')); return
        out_dir = os.path.join(DESKTOP_DIR, 'prism备份')
        os.makedirs(out_dir, exist_ok=True)
        html = ['<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8"><title>RT-QuIC Batch Report</title>',
                '<style>body{font-family:Arial,sans-serif;margin:24px;}h2{border-bottom:1px solid #ccc;padding-bottom:4px;}table{border-collapse:collapse;font-size:12px;margin:8px 0;}th,td{border:1px solid #bbb;padding:3px 7px;}th{background:#f0f0f0;}.num{text-align:right;}img{max-width:100%;border:1px solid #ddd;margin:6px 0;}</style></head><body>',
                f'<h1>RT-QuIC Batch Report</h1><p>Generated: {datetime.datetime.now().strftime("%Y-%m-%d %H:%M")} | Sheets: {len(sheets)}</p>']
        self._status.showMessage('⏳ 正在生成批量报告 ...')
        for name in sheets:
            html.append(f'<h2>📄 {name}</h2>')
            info = self._notes.get(name) or {}
            if info:
                html.append('<p><i>' + ' | '.join(f'{k}: {v}' for k, v in info.items() if v) + '</i></p>')
            rows = self._collect_sheet_metrics(name, self._last_metrics_threshold)
            if rows:
                keys = ['group', 'n', 'maxRFU', 'ttt', 't50', 'maxSlope', 'auc', 'lag', 'plateau']
                html.append('<table><tr>' + ''.join(f'<th>{k}</th>' for k in keys) + '</tr>')
                for r in rows:
                    html.append('<tr>' + ''.join(f'<td class="num">{r.get(k, "") if r.get(k) is not None else ""}</td>' for k in keys) + '</tr>')
                html.append('</table>')
            img = self._grab_view_image(name)
            if img:
                html.append(f'<img src="data:image/png;base64,{img}">')
        html.append('</body></html>')
        path = os.path.join(out_dir, 'RT-QuIC_Batch_Report.html')
        with open(path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(html))
        self._status.showMessage(f'✅ 批量报告: {path}')
        self._show_saved(path)

    def _tb_btn(self, text, cb, layout, icon=None, en=None):
        b = QPushButton(T(text, en or text)); b.setFont(FONTS.get('toolbar')); b.setCursor(Qt.PointingHandCursor)
        b.setStyleSheet('QPushButton{background:transparent;border:none;color:#1d1d1f;padding:5px 8px;border-radius:6px;}QPushButton:hover{background:#e8e8ed;}QPushButton:pressed{background:#dcdce1;}')
        if icon:
            b.setIcon(icons.make_icon(icon, '#1d1d1f', 15))
        b.clicked.connect(cb); layout.addWidget(b)
        self._tb_buttons.append((b, text, en or text))

    def _tb_sep(self, layout):
        s = QFrame(); s.setFrameShape(QFrame.VLine); s.setStyleSheet('QFrame{color:#d2d2d7;background:#d2d2d7;}')
        s.setFixedWidth(1); layout.addWidget(s)

    def _new_sheet(self, name=None, generate=True):
        if name is None: self._sheet_counter += 1; name = f'Sheet{self._sheet_counter}'
        base = name; i = 1
        while name in self._sheets: name = f'{base}_{i}'; i += 1
        self._sheets[name] = SheetModel()
        sw = SheetWidget(self._sheets[name], name); sw.edit_done.connect(self._auto_generate); sw.dirty_changed.connect(self._update_dirty_mark)
        sw.sync(); self._sheet_widgets[name] = sw
        self._tabs.addTab(sw, name); self._tabs.setCurrentWidget(sw)
        item = QTreeWidgetItem([f'  {name}']); item.setIcon(0, icons.make_icon('doc', '#6e6e73', 14)); item.setData(0, Qt.UserRole, name); item.setFont(0, FONTS.get('sidebar'))
        self._data_root.insertChild(self._data_root.childCount() - 1, item)
        self._data_root.setExpanded(True)
        self._tree.setCurrentItem(item)
        self._save_workspace(); self._status.showMessage(f'Created: {name}')
        self._refresh_all_tooltips()
        if generate:
            self._auto_generate(name)
        return name

    def _current_sheet_name(self):
        w = self._tabs.currentWidget()
        for n, sw in self._sheet_widgets.items():
            if sw is w: return n
        return None

    def _active_sheet_name(self):
        """当前数据表名；若当前选中的是视图，则返回该视图关联的数据表名。"""
        n = self._current_sheet_name()
        if n:
            return n
        w = self._tabs.currentWidget()
        for vn, vw in self._view_widgets.items():
            if vw is w:
                return self._view_sheet.get(vn)
        return None

    def _current_model(self):
        n = self._current_sheet_name()
        if n and n in self._sheet_widgets:
            self._sheet_widgets[n].model_from_table(); return self._sheets.get(n)
        return None

    def _update_dirty_mark(self, name):
        it = self._find_data_item(name)
        if not it: return
        sw = self._sheet_widgets.get(name)
        mark = ' •' if (sw and sw.is_dirty()) else ''
        it.setText(0, f'  {name}' + mark)

    def _update_item_tooltip(self, name, is_view=False):
        it = self._find_view_item(name) if is_view else self._find_data_item(name)
        if not it:
            return
        sheet = (self._view_sheet.get(name) or name) if is_view else name
        info = self._notes.get(sheet, {}) or {}
        text = '\n'.join(f'{k}: {v}' for k, v in info.items()) if info else T('无备注', 'No notes')
        it.setToolTip(0, text)

    def _refresh_all_tooltips(self):
        for i in range(self._data_root.childCount()):
            it = self._data_root.child(i)
            key = it.data(0, Qt.UserRole)
            if key and key != '__new__':
                self._update_item_tooltip(key, False)
        for i in range(self._view_root.childCount()):
            it = self._view_root.child(i)
            key = it.data(0, Qt.UserRole)
            if key and key != '__load_view__':
                self._update_item_tooltip(key, True)

    def _view_notes_action(self, vn):
        sheet = self._view_sheet.get(vn)
        if sheet and sheet in self._sheets:
            self._show_notes(sheet)
        else:
            QMessageBox.information(self, T('备注', 'Notes'), T('该视图未关联数据表。', 'This view is not linked to a sheet.'))

    def _undo(self):
        n = self._current_sheet_name()
        if n and n in self._sheet_widgets: self._sheet_widgets[n].undo()
    def _redo(self):
        n = self._current_sheet_name()
        if n and n in self._sheet_widgets: self._sheet_widgets[n].redo()

    def _ensure_plotly(self):
        views_dir = os.path.join(EDITOR_DIR, 'views'); os.makedirs(views_dir, exist_ok=True)
        pd = os.path.join(views_dir, 'plotly.min.js')
        if not os.path.exists(pd):
            try: shutil.copy2(PLOTLY_SRC, pd)
            except Exception: pass

    def _ensure_view_fresh(self, name):
        path = self._views.get(name)
        if not path or not os.path.exists(path):
            sn = self._view_sheet.get(name)
            if sn and sn in self._sheets: self._auto_generate(sn)
            return
        try:
            with open(path, encoding='utf-8') as f: html = f.read()
        except Exception:
            return
        if 'RTQUIC-ENGINE-V16' not in html:
            sn = self._view_sheet.get(name)
            if sn and sn in self._sheets: self._auto_generate(sn)

    def _update_status_info(self):
        if not hasattr(self, '_status_info'): return
        n = self._current_sheet_name()
        self._status_info.setText('当前: ' + (n or '—'))

    def _show_saved(self, paths):
        if isinstance(paths, str): paths = [paths]
        box = QMessageBox(self); box.setWindowTitle('导出完成')
        box.setText('已保存:\n' + '\n'.join(paths))
        reveal = box.addButton(T('在访达中显示', 'Show in Folder'), QMessageBox.ActionRole)
        box.addButton('确定', QMessageBox.AcceptRole)
        box.exec_()
        if box.clickedButton() == reveal and paths:
            _reveal_local_path(paths[0])

    def _setup_shortcuts(self):
        from PyQt5.QtWidgets import QShortcut
        from PyQt5.QtGui import QKeySequence
        QShortcut(QKeySequence.StandardKey.Open, self, self._import)
        QShortcut(QKeySequence.StandardKey.New, self, self._new_sheet)
        QShortcut(QKeySequence.StandardKey.Save, self, self._save_shortcut)
        QShortcut(QKeySequence.StandardKey.Undo, self, self._undo)
        QShortcut(QKeySequence.StandardKey.Redo, self, self._redo)
        QShortcut(QKeySequence('Ctrl+E'), self, self._export_csv)
        QShortcut(QKeySequence('Ctrl+Return'), self, self._generate_shortcut)

    def _save_shortcut(self):
        self._save_workspace(); self._status.showMessage('✅ 工作区已保存')

    def _generate_shortcut(self):
        n = self._current_sheet_name()
        if n: self._auto_generate(n)

    def _on_tree_clicked(self, item):
        key = item.data(0, Qt.UserRole)
        if key == '__new__':
            name, ok = QInputDialog.getText(self, 'New Sheet', 'Sheet name:')
            if ok and name.strip(): self._new_sheet(name.strip())
            elif ok: self._new_sheet()
        elif key == '__load_view__':
            self._load_view_file()
        elif key and self._is_data_item(item) and key in self._sheets:
            if key in self._sheet_widgets: self._tabs.setCurrentWidget(self._sheet_widgets[key])
            else:
                sw = SheetWidget(self._sheets[key], key); sw.edit_done.connect(self._auto_generate); sw.dirty_changed.connect(self._update_dirty_mark)
                sw.sync(); self._sheet_widgets[key] = sw
                self._tabs.addTab(sw, key); self._tabs.setCurrentWidget(sw)
        elif key and self._is_view_item(item) and key in self._views:
            self._ensure_view_fresh(key)
            if key in self._view_widgets: self._tabs.setCurrentWidget(self._view_widgets[key])
            else:
                vw = ViewWidget(self._views[key], view_name=key, window=self); self._view_widgets[key] = vw
                self._tabs.addTab(vw, f'📈 {key}'); self._tabs.setCurrentWidget(vw)
        self._update_status_info()

    def _is_data_item(self, item):
        p = item.parent()
        while p:
            if p is self._data_root: return True
            if p is self._view_root: return False
            p = p.parent()
        return False

    def _is_view_item(self, item):
        p = item.parent()
        while p:
            if p is self._view_root: return True
            if p is self._data_root: return False
            p = p.parent()
        return False

    def _find_data_item(self, key):
        for i in range(self._data_root.childCount()):
            it = self._data_root.child(i)
            if it.data(0, Qt.UserRole) == key:
                return it
        return None

    def _find_view_item(self, key):
        for i in range(self._view_root.childCount()):
            it = self._view_root.child(i)
            if it.data(0, Qt.UserRole) == key:
                return it
        return None

    def _reset_tree_fonts(self):
        if not hasattr(self, '_data_root') or not hasattr(self, '_view_root'):
            return
        f = FONTS.get('sidebar')
        for root in (self._data_root, self._view_root):
            root.setFont(0, f)
            for i in range(root.childCount()):
                root.child(i).setFont(0, f)

    def _sync_selection_bold(self, current, previous=None):
        """Bold the selected item and its corresponding Data/View item(s)."""
        if not hasattr(self, '_data_root') or not hasattr(self, '_view_root'):
            return
        self._reset_tree_fonts()
        if not current:
            return
        key = current.data(0, Qt.UserRole)
        if not key or key in ('__new__', '__load_view__'):
            return
        bold = [current]
        if self._is_data_item(current):
            for vn, sn in self._view_sheet.items():
                if sn == key:
                    it = self._find_view_item(vn)
                    if it:
                        bold.append(it)
        elif self._is_view_item(current):
            sn = self._view_sheet.get(key)
            if sn:
                it = self._find_data_item(sn)
                if it:
                    bold.append(it)
        bf = FONTS.get_bold('sidebar')
        for it in bold:
            it.setFont(0, bf)

    def _on_tree_drop(self):
        """After drag-drop, fix any cross-branch moves and pin special items."""
        for i in range(self._data_root.childCount() - 1, -1, -1):
            it = self._data_root.child(i)
            if it is self._new_item: continue
            if it.text(0).strip().startswith('📈'):
                self._data_root.removeChild(it)
                self._view_root.insertChild(self._view_root.childCount() - 1, it)
        for i in range(self._view_root.childCount() - 1, -1, -1):
            it = self._view_root.child(i)
            if it is self._load_item: continue
            if it.text(0).strip().startswith('📄'):
                self._view_root.removeChild(it)
                self._data_root.insertChild(self._data_root.childCount() - 1, it)
        # Pin "+ New Sheet" / "+ Load View" to the end
        self._data_root.removeChild(self._new_item)
        self._data_root.addChild(self._new_item)
        self._view_root.removeChild(self._load_item)
        self._view_root.addChild(self._load_item)
        self._save_workspace()

    def _on_tree_context_menu(self, pos):
        item = self._tree.itemAt(pos)
        if not item: return
        key = item.data(0, Qt.UserRole)
        if not key or key == '__new__': return
        menu = QMenu(self); is_data = self._is_data_item(item)
        if is_data: ra = menu.addAction('Rename'); da = menu.addAction('Delete')
        else:
            oa = menu.addAction('Open in Browser')
            na = menu.addAction(T('✏️ 编辑备注', 'Edit Notes'))
            pa = menu.addAction('导出PDF'); ta = menu.addAction('导出TIFF')
            ra = menu.addAction('Remove from View')
        action = menu.exec_(self._tree.mapToGlobal(pos))
        if not action: return
        if is_data:
            if action == ra:
                nn, ok = QInputDialog.getText(self, 'Rename', 'New name:', text=key)
                if ok and nn.strip() and nn.strip() != key:
                    nn = nn.strip(); self._sheets[nn] = self._sheets.pop(key)
                    if key in self._sheet_widgets:
                        self._sheet_widgets[nn] = self._sheet_widgets.pop(key)
                        idx = self._tabs.indexOf(self._sheet_widgets[nn])
                        if idx >= 0: self._tabs.setTabText(idx, nn)
                    item.setText(0, f'  {nn}'); item.setData(0, Qt.UserRole, nn)
                    # Rename the same-name view (auto-generated) to follow the data name
                    if key in self._views:
                        self._views[nn] = self._views.pop(key)
                        if key in self._view_widgets:
                            self._view_widgets[nn] = self._view_widgets.pop(key)
                            idx = self._tabs.indexOf(self._view_widgets[nn])
                            if idx >= 0: self._tabs.setTabText(idx, f'📈 {nn}')
                        if key in self._view_sheet:
                            self._view_sheet[nn] = self._view_sheet.pop(key)
                        vit = self._find_view_item(key)
                        if vit:
                            vit.setText(0, f'  {nn}'); vit.setData(0, Qt.UserRole, nn)
                    # Update the sheet link for any other (custom-named) views
                    for vn, sn in list(self._view_sheet.items()):
                        if sn == key: self._view_sheet[vn] = nn
                    if key in self._notes:
                        self._notes[nn] = self._notes.pop(key)
                    self._save_workspace()
                    self._refresh_all_tooltips()
            elif action == da:
                related_views = [vn for vn, sn in self._view_sheet.items() if sn == key]
                msg = f'Delete "{key}" and its data?'
                if related_views:
                    msg += f'\n\nThis will also delete {len(related_views)} linked view(s).'
                if QMessageBox.question(self, 'Delete', msg, QMessageBox.Yes | QMessageBox.No) == QMessageBox.Yes:
                    if key in self._sheet_widgets:
                        idx = self._tabs.indexOf(self._sheet_widgets[key])
                        if idx >= 0: self._tabs.removeTab(idx); del self._sheet_widgets[key]
                    del self._sheets[key]; self._data_root.removeChild(item)
                    for vn in related_views:
                        self._remove_view(vn, delete_file=True)
                    self._reset_tree_fonts(); self._save_workspace()
        else:
            if action == oa: _open_local_path(self._views[key])
            elif action == na: self._view_notes_action(key)
            elif action == pa: self._export_view_pdf(key)
            elif action == ta: self._export_view_tiff(key)
            elif action == ra:
                if QMessageBox.question(self, 'Delete View', f'Delete view "{key}"? (data is kept)', QMessageBox.Yes | QMessageBox.No) == QMessageBox.Yes:
                    self._remove_view(key, delete_file=True)
                    self._reset_tree_fonts()

    def _close_tab(self, idx):
        w = self._tabs.widget(idx)
        for n, sw in list(self._sheet_widgets.items()):
            if sw is w:
                if sw.is_dirty() and QMessageBox.question(self, '关闭', f'数据表「{n}」的改动尚未同步到视图，确定关闭吗？（数据已自动保存）', QMessageBox.Yes | QMessageBox.No) != QMessageBox.Yes:
                    return
                self._tabs.removeTab(idx); del self._sheet_widgets[n]; return
        for n, vw in list(self._view_widgets.items()):
            if vw is w: self._tabs.removeTab(idx); del self._view_widgets[n]; return
        if self._multiview_tab is not None and w is self._multiview_tab:
            self._tabs.removeTab(idx)
            self._multiview_widgets = []; self._multiview_tab = None
            return

    def _remove_view(self, name, delete_file=False):
        """Remove a view (tab, tree item, mapping) and optionally its HTML file."""
        if name in self._view_widgets:
            idx = self._tabs.indexOf(self._view_widgets[name])
            if idx >= 0:
                self._tabs.removeTab(idx)
            del self._view_widgets[name]
        path = self._views.pop(name, None)
        self._view_sheet.pop(name, None)
        it = self._find_view_item(name)
        if it:
            self._view_root.removeChild(it)
        if delete_file and path and os.path.exists(path):
            try:
                os.remove(path)
            except OSError:
                pass
        self._save_workspace()

    def _add(self, rows, cols):
        n = self._current_sheet_name()
        if not n: return
        m = self._sheets[n]; self._sheet_widgets[n].model_from_table()
        m.ensure(m.rows + rows, m.cols + cols); self._sheet_widgets[n].sync()

    def _clear(self):
        n = self._current_sheet_name()
        if not n: return
        if QMessageBox.question(self, T('清空', 'Clear'), T('清空当前表？', 'Clear current sheet?'), QMessageBox.Yes | QMessageBox.No) == QMessageBox.Yes:
            sw = self._sheet_widgets[n]
            sw.push_undo()
            self._sheets[n] = SheetModel(); sw.model = self._sheets[n]; sw.sync()
            sw.commit_change()

    def _import(self):
        path, _ = QFileDialog.getOpenFileName(self, T('导入', 'Import'), DESKTOP_DIR,
            'All Supported (*.csv *.tsv *.txt *.json *.xlsx);;Excel (*.xlsx);;CSV/TSV (*.csv);;JSON (*.json)')
        if not path: return
        self._import_one(path)

    def _import_one(self, path):
        try:
            fn = os.path.splitext(os.path.basename(path))[0]; name = self._new_sheet(fn, generate=False); m = self._sheets[name]
            if path.endswith('.xlsx'):
                wb = openpyxl.load_workbook(path, data_only=True); ws = wb.active
                matrix = [[str(c) if c is not None else '' for c in row]
                          for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=True)]
                wb.close(); m.from_text('\n'.join('\t'.join(row) for row in matrix))
            else:
                with open(path) as f: text = f.read()
                if path.endswith('.json'): m.from_json(text)
                else: m.from_text(text)
            self._sheet_widgets[name].sync(); self._save_history(fn)
            self._auto_generate(name)
        except Exception as e: QMessageBox.critical(self, T('导入错误', 'Import Error'), str(e))

    def _import_batch(self):
        paths, _ = QFileDialog.getOpenFileNames(self, T('批量导入', 'Batch Import'), DESKTOP_DIR,
            'All Supported (*.csv *.tsv *.txt *.json *.xlsx);;Excel (*.xlsx);;CSV/TSV (*.csv);;JSON (*.json)')
        if not paths: return
        for p in paths:
            self._import_one(p)
        self._status.showMessage(f'✅ 批量导入: {len(paths)} 个文件')

    def _export_csv(self):
        m = self._current_model()
        if not m: return
        path, _ = QFileDialog.getSaveFileName(self, 'Export CSV', 'data.csv', 'CSV (*.csv)')
        if path: open(path, 'w').write(m.to_csv())

    def _export_json(self):
        m = self._current_model()
        if not m: return
        path, _ = QFileDialog.getSaveFileName(self, 'Export JSON', 'data.json', 'JSON (*.json)')
        if path: open(path, 'w').write(m.to_json())

    def _aggregate_groups(self, model):
        """Group rows by identical sample name, compute mean per time point.
        Returns (aggregated_model, group_info) or (None, None) if no grouping needed."""
        groups = defaultdict(list)
        for i, name in enumerate(model.sample_names):
            name = name.strip()
            if not name:
                continue
            if i in (getattr(model, 'excluded', []) or []):
                continue
            row_vals = []
            for j in range(len(model.time_points)):
                v = model.get(i, j)
                row_vals.append(v if v is not None else np.nan)
            groups[name].append(row_vals)
        # Only keep groups with >1 replicate
        multi = {k: v for k, v in groups.items() if len(v) > 1}
        if not multi:
            return None, None
        info_parts = []
        new_model = SheetModel()
        new_model.time_points = list(model.time_points)
        new_model.sample_names = []
        new_model.data = []
        for name, rows in multi.items():
            arr = np.array(rows, dtype=float)
            means = np.nanmean(arr, axis=0)
            new_model.sample_names.append(name)
            new_model.data.append([None if np.isnan(m) else round(float(m), 2) for m in means])
            info_parts.append(f'{name} (n={len(rows)})')
        return new_model, '; '.join(info_parts)

    def _write_raw_sheet(self, wb, model):
        """Write per-replicate raw curves (Group, Excluded, values) for the viewer's 复孔 display."""
        try:
            time_cols = [j for j, tp in enumerate(model.time_points) if tp not in (None, '')]
            if not time_cols:
                return
            excl = getattr(model, 'excluded', []) or []
            ws = wb.create_sheet('RAW')
            ws.append(['Group', 'Excluded'] + [model.time_points[j] for j in time_cols])
            for i, s in enumerate(model.sample_names):
                if s is None or str(s).strip() == '':
                    continue
                row = [str(s).strip(), 1 if i in excl else 0]
                for j in time_cols:
                    v = model.get(i, j)
                    row.append(v if v is not None else '')
                ws.append(row)
        except Exception:
            import traceback; traceback.print_exc()

    def _write_notes_sheet(self, wb, name):
        """Write structured experiment notes (key, value) so the viewer shows them in the title."""
        try:
            info = self._notes.get(name) or {}
            if not info:
                return
            ws = wb.create_sheet('Notes')
            for k, v in info.items():
                if v:
                    ws.append([k, v])
        except Exception:
            pass

    def _add_sd_sheet(self, wb, model):
        """Add an 'SD' sheet to the workbook: per-group SD across replicates (for ±SD bands)."""
        try:
            time_cols = []
            for j, tp in enumerate(model.time_points):
                if tp is None or tp == '':
                    break
                time_cols.append(j)
            groups = defaultdict(list)
            for i, s in enumerate(model.sample_names):
                if s is None or s.strip() == '':
                    continue
                if i in (getattr(model, 'excluded', []) or []):
                    continue
                vals = [model.get(i, j) for j in time_cols]
                vals = [np.nan if x is None else x for x in vals]
                groups[s.strip()].append(vals)
            sd_rows = []
            for name, reps in groups.items():
                arr = np.array(reps, dtype=float)
                if arr.size == 0:
                    continue
                sds = np.nanstd(arr, axis=0)
                sd_rows.append([name] + [None if np.isnan(x) else round(float(x), 2) for x in sds])
            if not sd_rows:
                return
            ws2 = wb.create_sheet('SD')
            ws2.append([''] + [model.time_points[j] for j in time_cols])
            for row in sd_rows:
                ws2.append(row)
        except Exception:
            import traceback; traceback.print_exc()

    def _export_prism(self, name):
        """Export a sheet's MEAN curves (grouped by sample name) as a Prism XY CSV in ~/Desktop/RTedit→prism."""
        m = self._sheets.get(name)
        if not m:
            return None
        prism_dir = os.path.join(DESKTOP_DIR, 'RTedit→prism')
        os.makedirs(prism_dir, exist_ok=True)
        path = os.path.join(prism_dir, f'{name}.csv')
        # Valid time columns (contiguous from the start)
        time_cols = []
        for j, t in enumerate(m.time_points):
            if t is None or t == '':
                break
            time_cols.append(j)
        if not time_cols:
            return None
        # Group rows by sample name (replicates share the same name -> mean)
        groups = {}
        order = []
        for i, s in enumerate(m.sample_names):
            if s is None or s.strip() == '':
                continue
            sn = s.strip()
            row = [m.get(i, j) for j in time_cols]
            if sn not in groups:
                groups[sn] = []
                order.append(sn)
            groups[sn].append(row)
        if not order:
            return None
        rows = [['X'] + order]
        for k in range(len(time_cols)):
            line = [_to_hours(m.time_points[time_cols[k]])]
            for sn in order:
                vals = [g[k] for g in groups[sn] if g[k] is not None]
                line.append(round(sum(vals) / len(vals), 4) if vals else '')
            rows.append(line)
        with open(path, 'w', newline='') as f:
            csv.writer(f).writerows(rows)
        return path

    def _load_view_file(self):
        path, _ = QFileDialog.getOpenFileName(self, 'Load View File', DESKTOP_DIR,
            'HTML Files (*.html *.htm);;All Files (*)')
        if not path: return
        name = os.path.splitext(os.path.basename(path))[0]
        base = name; i = 1
        while name in self._views: name = f'{base}_{i}'; i += 1
        self._views[name] = path
        vw = ViewWidget(path, view_name=name, window=self); self._view_widgets[name] = vw
        self._tabs.addTab(vw, f'📈 {name}'); self._tabs.setCurrentWidget(vw)
        # Insert before the "+ Load View" item (last child)
        item = QTreeWidgetItem([f'  {name}']); item.setIcon(0, icons.make_icon('chartline', '#6e6e73', 14)); item.setData(0, Qt.UserRole, name); item.setFont(0, FONTS.get('sidebar'))
        self._view_root.insertChild(self._view_root.childCount() - 1, item)
        self._save_workspace(); self._status.showMessage(f'Loaded: {os.path.basename(path)}')
        self._refresh_all_tooltips()

    def _auto_generate(self, sheet_name):
        """Debounced auto-generate: create/update a same-name view for a sheet."""
        if self._building or sheet_name not in self._sheets: return
        m = self._sheets[sheet_name]
        if sheet_name in self._sheet_widgets:
            self._sheet_widgets[sheet_name].model_from_table()
        self._export_prism(sheet_name)
        agg, info = self._aggregate_groups(m)
        if agg: m = agg
        views_dir = os.path.join(EDITOR_DIR, 'views')
        os.makedirs(views_dir, exist_ok=True)
        csv_text = m.to_csv()
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sheet1'
        for row in csv.reader(io.StringIO(csv_text)):
            ws.append([None if c.strip() == '' else (float(c) if c.strip().replace('.', '').replace('-', '').isdigit() else c) for c in row])
        self._add_sd_sheet(wb, self._sheets[sheet_name])
        self._write_raw_sheet(wb, self._sheets[sheet_name])
        self._write_notes_sheet(wb, sheet_name)
        xlsx_path = os.path.join(views_dir, f'{sheet_name}.xlsx'); wb.save(xlsx_path)
        self._building = True

        def run():
            try:
                bv.build(xlsx_path, output_dir=views_dir, open_browser=False)
                generated = os.path.join(views_dir, f'{sheet_name}_viewer.html')
                if os.path.exists(generated):
                    with open(generated, 'r') as f: html = f.read()
                    html = html.replace('Times New Roman', 'Arial')
                    with open(generated, 'w') as f: f.write(html)
                pd = os.path.join(views_dir, 'plotly.min.js')
                if not os.path.exists(pd): shutil.copy2(PLOTLY_SRC, pd)
                self.auto_gen_done_signal.emit(generated, sheet_name, sheet_name)
            except Exception:
                self.auto_gen_done_signal.emit('', sheet_name, sheet_name)
        threading.Thread(target=run, daemon=True).start()

    def _on_auto_gen_done(self, out_path, out_name, sheet_name):
        self._building = False
        if not out_path or not os.path.exists(out_path): return
        if out_name in self._views:
            self._views[out_name] = out_path
            self._view_sheet[out_name] = sheet_name
            if out_name in self._view_widgets:
                self._view_widgets[out_name].reload()
        else:
            if out_name in self._view_widgets:
                self._view_widgets[out_name].reload()
            else:
                vw = ViewWidget(out_path, view_name=out_name, window=self); self._view_widgets[out_name] = vw
                self._views[out_name] = out_path; self._view_sheet[out_name] = sheet_name
                self._tabs.addTab(vw, f'📈 {out_name}')
            if not self._find_view_item(out_name):
                item = QTreeWidgetItem([f'  {out_name}']); item.setIcon(0, icons.make_icon('chartline', '#6e6e73', 14)); item.setData(0, Qt.UserRole, out_name); item.setFont(0, FONTS.get('sidebar'))
                self._view_root.insertChild(self._view_root.childCount() - 1, item)
        self._save_workspace()
        sw = self._sheet_widgets.get(sheet_name)
        if sw: sw.mark_clean()
        self._update_dirty_mark(sheet_name)
        self._status.showMessage(f'✅ Auto-updated view: {out_name}')
        self._refresh_all_tooltips()

    def _refresh_all_views(self):
        """Regenerate every sheet's view on startup so HTML always matches the engine."""
        sheets = list(self._sheets.keys())
        if not sheets:
            return
        views_dir = os.path.join(EDITOR_DIR, 'views')
        os.makedirs(views_dir, exist_ok=True)

        # Ensure plotly.min.js is present so views can render
        pd = os.path.join(views_dir, 'plotly.min.js')
        if not os.path.exists(pd):
            try:
                shutil.copy2(PLOTLY_SRC, pd)
            except Exception:
                pass

        def run():
            for name in sheets:
                try:
                    if name not in self._sheets:
                        continue
                    m = self._sheets[name]
                    self._export_prism(name)
                    agg, info = self._aggregate_groups(m)
                    if agg:
                        m = agg
                    csv_text = m.to_csv()
                    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sheet1'
                    for row in csv.reader(io.StringIO(csv_text)):
                        ws.append([None if c.strip() == '' else (float(c) if c.strip().replace('.', '').replace('-', '').isdigit() else c) for c in row])
                    self._add_sd_sheet(wb, self._sheets[name])
                    self._write_raw_sheet(wb, self._sheets[name])
                    self._write_notes_sheet(wb, name)
                    xlsx_path = os.path.join(views_dir, f'{name}.xlsx'); wb.save(xlsx_path)
                    bv.build(xlsx_path, output_dir=views_dir, open_browser=False)
                    generated = os.path.join(views_dir, f'{name}_viewer.html')
                    if os.path.exists(generated):
                        with open(generated, 'r') as f: html = f.read()
                        html = html.replace('Times New Roman', 'Arial')
                        with open(generated, 'w') as f: f.write(html)
                    if not os.path.exists(pd):
                        shutil.copy2(PLOTLY_SRC, pd)
                    self.auto_gen_done_signal.emit(generated, name, name)
                except Exception:
                    self.auto_gen_done_signal.emit('', name, name)

        threading.Thread(target=run, daemon=True).start()


    def _metrics(self, values, times, threshold=1000):
        """Compute RT-QuIC kinetics metrics for one curve (mirrors viewer calcM + plateau)."""
        v = [float(x) if x is not None else 0.0 for x in values]
        t = [float(x) if x is not None else 0.0 for x in times]
        if not v:
            return None
        mx = max(v); mxi = v.index(mx)
        ttt = None
        for i in range(len(v)):
            if v[i] >= threshold:
                ttt = t[i]; break
        t50 = None; half = mx / 2
        for i in range(len(v)):
            if v[i] >= half:
                if i > 0 and v[i - 1] < half and v[i] != v[i - 1]:
                    t50 = t[i - 1] + (t[i] - t[i - 1]) * (half - v[i - 1]) / (v[i] - v[i - 1])
                else:
                    t50 = t[i]
                break
        ms = 0.0; mst = t[0] if t else 0; msv = v[0] if v else 0
        for i in range(len(v) - 5):
            dx = t[i + 5] - t[i]
            if dx > 0:
                s = (v[i + 5] - v[i]) / dx
                if s > ms:
                    ms = s; mst = t[i + 2] if i + 2 < len(t) else t[-1]; msv = v[i + 2] if i + 2 < len(v) else v[-1]
        bg = sum(v[:5]) / 5 if len(v) >= 5 else (sum(v) / len(v) if v else 0.0)
        mpr = mx / bg if bg > 0 else 0.0
        auc = 0.0
        for i in range(1, len(v)):
            y0 = max(0.0, v[i - 1] - bg); y1 = max(0.0, v[i] - bg)
            auc += (y0 + y1) / 2 * (t[i] - t[i - 1])
        lag = None
        if ms > 0 and msv > bg:
            lg = mst - (msv - bg) / ms
            if 0 <= lg <= mst:
                lag = lg
        plt = None
        pv = [x for x in v if x >= 0.95 * mx]
        if pv:
            plt = sum(pv) / len(pv)
        return {'maxRFU': mx, 'maxTime': t[mxi], 'ttt': ttt, 't50': t50, 'maxSlope': ms,
                'bg': bg, 'mpr': mpr, 'auc': auc, 'lag': lag, 'plateau': plt}

    def _collect_sheet_metrics(self, sheet_name, threshold=1000):
        """Per-group mean metrics + replicate CV% for one sheet."""
        m = self._sheets.get(sheet_name)
        if not m:
            return []
        time_cols = []
        for j, tp in enumerate(m.time_points):
            if tp is None or tp == '':
                break
            time_cols.append(j)
        if not time_cols:
            return []
        times = []
        for j in time_cols:
            try:
                times.append(float(m.time_points[j]))
            except (ValueError, TypeError):
                times.append(0.0)
        groups = defaultdict(list)
        for i, s in enumerate(m.sample_names):
            if s is None or s.strip() == '':
                continue
            if i in (getattr(m, 'excluded', []) or []):
                continue
            vals = [m.get(i, j) for j in time_cols]
            vals = [np.nan if x is None else x for x in vals]
            groups[s.strip()].append(vals)
        keys = ['maxRFU', 'ttt', 't50', 'maxSlope', 'auc', 'lag', 'plateau']
        rows = []
        for name, reps in groups.items():
            met = [mm for r in reps if (mm := self._metrics(r, times, threshold))]
            out = {'group': name, 'n': len(met)}
            for k in keys:
                vals = [x[k] for x in met if x[k] is not None]
                if vals:
                    meanv = sum(vals) / len(vals)
                    out[k] = round(meanv, 2)
                    if len(vals) > 1 and meanv != 0:
                        sd = (sum((x - meanv) ** 2 for x in vals) / len(vals)) ** 0.5
                        out[k + '_CV%'] = round(abs(sd / meanv * 100), 1)
                    else:
                        out[k + '_CV%'] = 0.0
                else:
                    out[k] = None
                    out[k + '_CV%'] = None
            rows.append(out)
        return rows

    def _export_metrics(self):
        """Export current sheet's per-group metrics (mean + CV%) to CSV & XLSX in ~/Desktop/RTedit→prism."""
        n = self._current_sheet_name()
        if not n:
            QMessageBox.information(self, 'Export Metrics', 'Please open a Data sheet first.'); return
        threshold, ok = QInputDialog.getInt(self, 'TtT Threshold', 'Threshold (RFU):', self._last_metrics_threshold, 1, 10000000)
        if not ok:
            return
        self._last_metrics_threshold = threshold
        rows = self._collect_sheet_metrics(n, threshold)
        if not rows:
            QMessageBox.information(self, 'Export Metrics', 'No data to export.'); return
        out_dir = os.path.join(DESKTOP_DIR, 'RTedit→prism')
        os.makedirs(out_dir, exist_ok=True)
        keys = ['group', 'n', 'maxRFU', 'maxRFU_CV%', 'ttt', 'ttt_CV%', 't50', 't50_CV%',
                'maxSlope', 'maxSlope_CV%', 'auc', 'auc_CV%', 'lag', 'lag_CV%', 'plateau', 'plateau_CV%']
        csv_path = os.path.join(out_dir, f'{n}_metrics.csv')
        with open(csv_path, 'w', newline='') as f:
            w = csv.writer(f); w.writerow(keys)
            for r in rows:
                w.writerow([r.get(k, '') if r.get(k) is not None else '' for k in keys])
        xlsx_path = os.path.join(out_dir, f'{n}_metrics.xlsx')
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Metrics'
        ws.append(keys)
        for r in rows:
            ws.append([r.get(k, '') if r.get(k) is not None else '' for k in keys])
        wb.save(xlsx_path)
        self._status.showMessage(f'✅ Metrics exported: {csv_path}')
        self._show_saved([csv_path, xlsx_path])

    def _show_notes(self, n=None):
        """自由填写实验信息：字段名和值都可自由输入（文本/数值均可），支持动态增删行。

        注意：工具栏按钮的 clicked 信号会把 checked=False 传给 n，因此这里必须用
        真值判断（而不是 `is None`），把 False 与 None 一样按"未指定"处理。
        """
        if not n:
            n = self._active_sheet_name()
            if not n and self._tabs.currentWidget() is getattr(self, '_multiview_tab', None):
                # 多视图页签：从同屏视图中选择要编辑备注的视图（仅一个时直接打开）
                names = [vn for vn in getattr(self, '_multiview_names', []) if self._view_sheet.get(vn)]
                if not names:
                    pass
                elif len(names) == 1:
                    n = self._view_sheet.get(names[0])
                else:
                    vn, ok = QInputDialog.getItem(self, T('备注', 'Notes'),
                                                  T('选择要编辑备注的视图', 'Choose a view to edit notes'),
                                                  names, 0, False)
                    if ok: n = self._view_sheet.get(vn)
        if not n:
            QMessageBox.information(self, T('备注', 'Notes'), T('请先打开一个数据表。', 'Please open a Data sheet first.')); return
        info = self._notes.get(n, {}) or {}
        dlg = QDialog(self); dlg.setWindowTitle(f'{T("实验信息", "Notes")} — {n}'); dlg.setMinimumSize(500, 460)
        lay = QVBoxLayout(dlg)
        lay.addWidget(QLabel(T('字段名和值都可自由填写（文本/数值均可），一行一个字段：', 'Field name and value are both free-form (text or number), one field per row:')))
        table = QTableWidget(); table.setColumnCount(2)
        table.setHorizontalHeaderLabels([T('字段名', 'Field'), T('值', 'Value')])
        table.horizontalHeader().setStretchLastSection(True)
        table.setColumnWidth(0, 180)
        items = list(info.items()) or [('', '')]
        table.setRowCount(len(items))
        for i, (k, v) in enumerate(items):
            table.setItem(i, 0, QTableWidgetItem(str(k)))
            table.setItem(i, 1, QTableWidgetItem(str(v)))
        lay.addWidget(table)
        btns_row = QHBoxLayout()
        add_btn = QPushButton(T('➕ 添加字段', '+ Add field'))
        del_btn = QPushButton(T('🗑 删除所选行', 'Delete selected row'))
        tpl_load_btn = QPushButton(T('📄 载入模板', '📄 Load Template'))
        tpl_load_btn.setToolTip(T('载入已保存的备注模板，载入后可直接修改', 'Load a saved notes template, then edit it directly'))
        tpl_save_btn = QPushButton(T('💾 存为模板', '💾 Save as Template'))
        tpl_save_btn.setToolTip(T('把当前备注字段保存为模板，供之后新建备注时载入', 'Save the current fields as a template for later notes'))
        btns_row.addWidget(add_btn); btns_row.addWidget(del_btn)
        btns_row.addStretch()
        btns_row.addWidget(tpl_load_btn); btns_row.addWidget(tpl_save_btn)
        lay.addLayout(btns_row)
        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.accepted.connect(dlg.accept); btns.rejected.connect(dlg.reject)
        lay.addWidget(btns)

        def _table_rows():
            rows = []
            for i in range(table.rowCount()):
                it0 = table.item(i, 0); it1 = table.item(i, 1)
                k = it0.text().strip() if it0 else ''
                v = it1.text().strip() if it1 else ''
                if k:
                    rows.append([k, v])
            return rows

        def _fill_rows(rows):
            """用模板内容填充表格（替换当前行），填充后可直接修改。"""
            table.setRowCount(0)
            for k, v in rows:
                r = table.rowCount()
                table.insertRow(r)
                table.setItem(r, 0, QTableWidgetItem(str(k)))
                table.setItem(r, 1, QTableWidgetItem(str(v)))

        def add_row():
            r = table.rowCount()
            table.insertRow(r)
            table.setItem(r, 0, QTableWidgetItem(''))
            table.setItem(r, 1, QTableWidgetItem(''))
            table.setCurrentCell(r, 0)
        def del_row():
            r = table.currentRow()
            if r >= 0:
                table.removeRow(r)

        def save_tpl():
            rows = _table_rows()
            if not rows:
                QMessageBox.information(dlg, T('存为模板', 'Save as Template'),
                                        T('当前备注为空，请先填写字段再保存模板。', 'Notes are empty; add fields before saving a template.'))
                return
            default_name = datetime.datetime.now().strftime(T('模板%m%d%H%M', 'Template%m%d%H%M'))
            name, ok = QInputDialog.getText(dlg, T('存为模板', 'Save as Template'),
                                            T('模板名称：', 'Template name:'), text=default_name)
            if not ok or not name.strip():
                return
            tpls = load_note_templates()
            tpls[name.strip()] = rows
            save_note_templates(tpls)
            self._status.showMessage(f'✅ {T("模板已保存", "Template saved")}: {name.strip()}')

        def load_tpl():
            tpls = load_note_templates()
            if not tpls:
                QMessageBox.information(dlg, T('载入模板', 'Load Template'),
                                        T('还没有保存过模板，先用"💾 存为模板"保存一个。', 'No saved templates yet; use "💾 Save as Template" first.'))
                return
            m = QMenu(dlg)
            for name in tpls:
                m.addAction(name, lambda nm=name: _fill_rows(tpls[nm]))
            m.addSeparator()
            m.addAction(T('🗑 删除模板...', '🗑 Delete Template...'), lambda: del_tpl_menu(tpls))
            m.exec_(tpl_load_btn.mapToGlobal(tpl_load_btn.rect().bottomLeft()))

        def del_tpl_menu(tpls):
            m2 = QMenu(dlg)
            def remove(nm):
                tpls.pop(nm, None)
                save_note_templates(tpls)
                self._status.showMessage(f'🗑 {T("模板已删除", "Template deleted")}: {nm}')
            for name in list(tpls):
                m2.addAction(name, lambda nm=name: remove(nm))
            m2.exec_(tpl_load_btn.mapToGlobal(tpl_load_btn.rect().bottomLeft()))

        add_btn.clicked.connect(add_row)
        del_btn.clicked.connect(del_row)
        tpl_load_btn.clicked.connect(load_tpl)
        tpl_save_btn.clicked.connect(save_tpl)

        if dlg.exec_() == QDialog.Accepted:
            nd = {}
            for i in range(table.rowCount()):
                it0 = table.item(i, 0); it1 = table.item(i, 1)
                k = it0.text().strip() if it0 else ''
                v = it1.text().strip() if it1 else ''
                if k:
                    nd[k] = v
            self._notes[n] = nd
            self._save_workspace()
            self._status.showMessage(f'✅ 实验信息已保存: {n}')
            self._refresh_all_tooltips()

    def _export_report(self):
        """Generate an HTML summary report (all sheets: notes + metrics + view links)."""
        threshold, ok = QInputDialog.getInt(self, 'TtT Threshold', 'Threshold (RFU):', 1000, 1, 10000000)
        if not ok:
            return
        html = ['<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8"><title>RT-QuIC Report</title>',
                '<style>body{font-family:Arial,sans-serif;margin:24px;color:#000;}h1{font-size:22px;}h2{font-size:17px;border-bottom:1px solid #ccc;padding-bottom:4px;}table{border-collapse:collapse;font-size:12px;margin:8px 0;}th,td{border:1px solid #bbb;padding:4px 8px;}th{background:#f0f0f0;}.num{text-align:right;}</style></head><body>',
                f'<h1>RT-QuIC Summary Report</h1><p>Generated: {datetime.datetime.now().strftime("%Y-%m-%d %H:%M")}</p>']
        for name in self._sheets:
            notes = self._notes.get(name, {}).get('note', '')
            html.append(f'<h2>📄 {name}</h2>')
            if notes:
                html.append(f'<p><b>Notes:</b> {notes.replace(chr(10), "<br>")}</p>')
            rows = self._collect_sheet_metrics(name, threshold)
            if rows:
                keys = ['group', 'n', 'maxRFU', 'ttt', 't50', 'maxSlope', 'auc', 'lag', 'plateau']
                html.append('<table><tr>' + ''.join(f'<th>{k}</th>' for k in keys) + '</tr>')
                for r in rows:
                    html.append('<tr>' + ''.join(f'<td class="num">{r.get(k, "") if r.get(k) is not None else ""}</td>' for k in keys) + '</tr>')
                html.append('</table>')
            else:
                html.append('<p><i>No metrics data.</i></p>')
            views_here = [vn for vn, sn in self._view_sheet.items() if sn == name]
            if views_here:
                links = ' | '.join(f'<a href="file://{self._views[vn]}">{vn}</a>' for vn in views_here if vn in self._views)
                html.append(f'<p>Views: {links}</p>')
            else:
                html.append('<p>No view yet.</p>')
        html.append('</body></html>')
        out_dir = os.path.join(DESKTOP_DIR, 'prism备份')
        os.makedirs(out_dir, exist_ok=True)
        path = os.path.join(out_dir, 'RT-QuIC_Report.html')
        with open(path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(html))
        _open_local_path(path)
        self._status.showMessage(f'✅ Report: {path}')

    def _save_download(self, item):
        try:
            out_dir = os.path.join(DESKTOP_DIR, 'prism备份')
            os.makedirs(out_dir, exist_ok=True)
            fn = item.suggestedFileName() or 'rtquic_download'
            path = os.path.join(out_dir, fn)
            item.setPath(path)
            item.accept()
            self._status.showMessage(f'✅ 导出完成: {os.path.basename(path)}', 4000)
        except Exception:
            item.cancel()

    def _export_view_pdf(self, name):
        vw = self._view_widgets.get(name)
        if not vw or not getattr(vw, '_web', None):
            return
        out_dir = os.path.join(DESKTOP_DIR, 'prism备份')
        os.makedirs(out_dir, exist_ok=True)
        path = os.path.join(out_dir, f'{name}.pdf')
        try:
            def _pdf_done(success):
                if success:
                    self._status.showMessage('✅ 导出完成: ' + os.path.basename(path), 4000)
                    self._show_saved(path)
                else:
                    self._status.showMessage('❌ PDF 导出失败', 4000)
            vw._web.page().printToPdf(path).finished.connect(_pdf_done)
            self._status.showMessage('正在导出 PDF ...', 2000)
        except Exception as e:
            QMessageBox.information(self, 'Export PDF', f'Failed: {e}')

    def _export_view_tiff(self, name):
        vw = self._view_widgets.get(name)
        if not vw or not getattr(vw, '_web', None):
            return
        out_dir = os.path.join(DESKTOP_DIR, 'prism备份')
        os.makedirs(out_dir, exist_ok=True)
        path = os.path.join(out_dir, f'{name}.tiff')
        pix = vw._web.grab()
        if not pix.isNull():
            pix.save(path, 'TIFF')
            self._status.showMessage('✅ 导出完成: ' + os.path.basename(path), 4000)
            self._show_saved(path)
        else:
            self._status.showMessage('❌ TIFF 导出失败', 4000)

    def _generate(self):
        m = self._current_model()
        if not m or self._building: return
        n0 = self._current_sheet_name() or 'rt-quic-viewer'
        default = os.path.join(DESKTOP_DIR, n0 + '.html')
        out_path, _ = QFileDialog.getSaveFileName(self, T('保存可视化', 'Save Visualization'), default, 'HTML (*.html)')
        if not out_path: return
        out_dir = os.path.dirname(out_path); out_name = os.path.splitext(os.path.basename(out_path))[0]

        # Auto-detect replicate groups and compute means
        agg_model, group_info = self._aggregate_groups(m)
        if agg_model:
            m = agg_model
            self._status.showMessage(f'🔬 Auto-grouped: {group_info}')

        csv_text = m.to_csv()
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sheet1'
        for row in csv.reader(io.StringIO(csv_text)):
            ws.append([None if c.strip() == '' else (float(c) if c.strip().replace('.', '').replace('-', '').isdigit() else c) for c in row])
        cur = self._current_sheet_name()
        if cur and cur in self._sheets:
            self._add_sd_sheet(wb, self._sheets[cur])
            self._write_raw_sheet(wb, self._sheets[cur])
            self._write_notes_sheet(wb, cur)
        xlsx_path = os.path.join(EDITOR_DIR, f'{out_name}.xlsx'); wb.save(xlsx_path)
        self._gen_btn.setEnabled(False); self._building = True

        def run():
            try:
                bv.build(xlsx_path, output_dir=out_dir, open_browser=False)
                generated = os.path.join(out_dir, f'{out_name}_viewer.html')
                if os.path.exists(generated) and generated != out_path: shutil.move(generated, out_path)
                if os.path.exists(out_path):
                    with open(out_path, 'r') as f: html = f.read()
                    html = html.replace('Times New Roman', 'Arial')
                    with open(out_path, 'w') as f: f.write(html)
                pd = os.path.join(out_dir, 'plotly.min.js')
                if not os.path.exists(pd): shutil.copy2(PLOTLY_SRC, pd)
                self.gen_done_signal.emit(out_path, out_name)
            except Exception as e: self.gen_done_signal.emit('', str(e))
        threading.Thread(target=run, daemon=True).start()

    def _on_gen_done(self, out_path, out_name):
        self._building = False; self._gen_btn.setEnabled(True)
        if not out_path: QMessageBox.critical(self, 'Error', out_name); return
        sheet = self._current_sheet_name()
        vw = ViewWidget(out_path, view_name=out_name, window=self); self._view_widgets[out_name] = vw; self._views[out_name] = out_path
        self._view_sheet[out_name] = sheet
        self._tabs.addTab(vw, f'📈 {out_name}'); self._tabs.setCurrentWidget(vw)
        item = QTreeWidgetItem([f'  {out_name}']); item.setIcon(0, icons.make_icon('chartline', '#6e6e73', 14)); item.setData(0, Qt.UserRole, out_name); item.setFont(0, FONTS.get('sidebar'))
        self._view_root.insertChild(self._view_root.childCount() - 1, item); self._view_root.setExpanded(True)
        self._save_workspace(); self._status.showMessage(f'✅ Generated: {os.path.basename(out_path)}')
        self._refresh_all_tooltips()

    def _workspace_path(self):
        return os.path.join(EDITOR_DIR, 'workspace.json')

    def _save_workspace(self):
        data_order = [self._data_root.child(i).data(0, Qt.UserRole)
                      for i in range(self._data_root.childCount())
                      if self._data_root.child(i).data(0, Qt.UserRole) != '__new__']
        view_order = [self._view_root.child(i).data(0, Qt.UserRole)
                      for i in range(self._view_root.childCount())
                      if self._view_root.child(i).data(0, Qt.UserRole) != '__load_view__']
        ws = {'sheets': {}, 'views': {vn: {'path': p, 'sheet': self._view_sheet.get(vn)} for vn, p in self._views.items()},
              'sheet_counter': self._sheet_counter, 'data_order': data_order, 'view_order': view_order, 'notes': self._notes}
        for name, sw in self._sheet_widgets.items(): sw.model_from_table()
        for name, m in self._sheets.items():
            ws['sheets'][name] = {'timePoints': m.time_points, 'sampleNames': m.sample_names, 'data': m.data,
                                 'excluded': getattr(m, 'excluded', [])}
        with open(self._workspace_path(), 'w') as f: json.dump(ws, f, ensure_ascii=False)

    def _load_workspace(self):
        path = self._workspace_path()
        if not os.path.exists(path): return
        try:
            with open(path) as f: ws = json.load(f)
            self._sheet_counter = ws.get('sheet_counter', 0)
            self._notes = ws.get('notes', {})
            sheets = ws.get('sheets', {})
            views = ws.get('views', {})
            data_order = ws.get('data_order') or list(sheets.keys())
            view_order = ws.get('view_order') or list(views.keys())
            for name in data_order:
                if name not in sheets: continue
                d = sheets[name]
                m = SheetModel(); m.time_points = d.get('timePoints', []); m.sample_names = d.get('sampleNames', []); m.data = d.get('data', []); m.excluded = list(d.get('excluded', []))
                self._sheets[name] = m
                sw = SheetWidget(m, name); sw.edit_done.connect(self._auto_generate); sw.dirty_changed.connect(self._update_dirty_mark); sw.sync(); self._sheet_widgets[name] = sw
                self._tabs.addTab(sw, name)
                item = QTreeWidgetItem([f'  {name}']); item.setIcon(0, icons.make_icon('doc', '#6e6e73', 14)); item.setData(0, Qt.UserRole, name); item.setFont(0, FONTS.get('sidebar'))
                self._data_root.insertChild(self._data_root.childCount() - 1, item)
            for name in sheets:
                if name not in data_order:
                    d = sheets[name]
                    m = SheetModel(); m.time_points = d.get('timePoints', []); m.sample_names = d.get('sampleNames', []); m.data = d.get('data', []); m.excluded = list(d.get('excluded', []))
                    self._sheets[name] = m
                    sw = SheetWidget(m, name); sw.edit_done.connect(self._auto_generate); sw.dirty_changed.connect(self._update_dirty_mark); sw.sync(); self._sheet_widgets[name] = sw
                    self._tabs.addTab(sw, name)
                    item = QTreeWidgetItem([f'  {name}']); item.setIcon(0, icons.make_icon('doc', '#6e6e73', 14)); item.setData(0, Qt.UserRole, name); item.setFont(0, FONTS.get('sidebar'))
                    self._data_root.insertChild(self._data_root.childCount() - 1, item)
            for vn in view_order:
                if vn not in views: continue
                vd = views[vn]
                vpath = vd.get('path', '')
                if os.path.exists(vpath):
                    self._views[vn] = vpath; self._view_sheet[vn] = vd.get('sheet')
                    item = QTreeWidgetItem([f'  {vn}']); item.setIcon(0, icons.make_icon('chartline', '#6e6e73', 14)); item.setData(0, Qt.UserRole, vn); item.setFont(0, FONTS.get('sidebar'))
                    self._view_root.insertChild(self._view_root.childCount() - 1, item)
            for vn, vd in views.items():
                if vn not in view_order:
                    vpath = vd.get('path', '')
                    if os.path.exists(vpath):
                        self._views[vn] = vpath; self._view_sheet[vn] = vd.get('sheet')
                        item = QTreeWidgetItem([f'  {vn}']); item.setIcon(0, icons.make_icon('chartline', '#6e6e73', 14)); item.setData(0, Qt.UserRole, vn); item.setFont(0, FONTS.get('sidebar'))
                        self._view_root.insertChild(self._view_root.childCount() - 1, item)
            self._status.showMessage('Workspace restored')
            self._refresh_all_tooltips()
        except Exception: self._status.showMessage('⚠ 工作区恢复失败')

    def closeEvent(self, event):
        dirty = [n for n, sw in self._sheet_widgets.items() if sw.is_dirty()]
        if dirty:
            names = '、'.join(dirty[:5]) + (' 等' if len(dirty) > 5 else '')
            if QMessageBox.question(self, '退出', f'有 {len(dirty)} 个数据表的改动尚未同步到视图（数据已自动保存）。确定退出？\n{names}', QMessageBox.Yes | QMessageBox.No) != QMessageBox.Yes:
                event.ignore(); return
        try:
            s = load_settings(); s['win_geo'] = [self.x(), self.y(), self.width(), self.height()]; save_settings(s)
        except Exception:
            pass
        self._save_workspace(); super().closeEvent(event)

    def _save_history(self, label=None):
        hist_dir = os.path.join(EDITOR_DIR, 'history'); os.makedirs(hist_dir, exist_ok=True)
        ts = datetime.datetime.now().strftime('%Y-%m-%d_%H-%M-%S')
        m = self._current_model()
        if not m: return
        name = label or (m.sample_names[0] or 'untitled')[:20]
        name = ''.join(c for c in name if c.isalnum() or c in ' _.-') or 'untitled'
        open(os.path.join(hist_dir, f'{ts}_{name}.json'), 'w').write(m.to_json())

    def _show_history(self):
        hist_dir = os.path.join(EDITOR_DIR, 'history')
        if not os.path.exists(hist_dir): QMessageBox.information(self, 'History', 'No history yet.'); return
        files = sorted([f for f in os.listdir(hist_dir) if f.endswith('.json')], reverse=True)
        if not files: QMessageBox.information(self, 'History', 'No history yet.'); return
        from PyQt5.QtWidgets import QDialog, QListWidget, QDialogButtonBox, QVBoxLayout as VBox
        dlg = QDialog(self); dlg.setWindowTitle('History'); dlg.setMinimumSize(500, 350)
        layout = VBox(dlg); lst = QListWidget()
        items = []
        for f in files:
            parts = f.replace('.json', '').split('_', 2)
            label = f'{parts[0]} {parts[1].replace("-", ":")}  —  {parts[2]}' if len(parts) >= 3 else f
            items.append(label); lst.addItem(label)
        layout.addWidget(lst)
        btns = QDialogButtonBox()
        lb = btns.addButton('Load', QDialogButtonBox.AcceptRole)
        db = btns.addButton('Delete', QDialogButtonBox.DestructiveRole)
        btns.addButton('Cancel', QDialogButtonBox.RejectRole); layout.addWidget(btns)

        def click(btn):
            idx = lst.currentRow()
            if btn == db and idx >= 0:
                if QMessageBox.question(dlg, 'Delete', f'Delete "{items[idx]}"?', QMessageBox.Yes | QMessageBox.No) == QMessageBox.Yes:
                    os.remove(os.path.join(hist_dir, files[idx])); lst.takeItem(idx); del items[idx], files[idx]
                    if lst.count() == 0: dlg.close()
            elif btn == lb and idx >= 0: dlg.accept()
        btns.clicked.connect(click)

        if dlg.exec_() == QDialog.Accepted and lst.currentRow() >= 0:
            idx = lst.currentRow()
            d = json.load(open(os.path.join(hist_dir, files[idx])))
            m = self._current_model()
            if m:
                m.time_points = d.get('timePoints', []); m.sample_names = d.get('sampleNames', []); m.data = d.get('data', [])
                n = self._current_sheet_name()
                if n: self._sheet_widgets[n].sync()
                self._status.showMessage(f'Loaded: {files[idx]}')

    def keyPressEvent(self, event):
        if event.modifiers() == Qt.ControlModifier and event.key() == Qt.Key_V:
            text = QApplication.clipboard().text()
            if text:
                lines = [l for l in text.strip().split('\n') if l.strip()]
                if len(lines) >= 2:
                    n = self._current_sheet_name()
                    m = self._current_model()
                    if m and n and n in self._sheet_widgets:
                        sw = self._sheet_widgets[n]
                        sw.push_undo()
                        m.from_text(text)
                        sw.sync()
                        sw.commit_change()
                        self._save_history('pasted')
                        self._save_workspace()
        else: super().keyPressEvent(event)


# ============================== APPLE STYLE QSS ===============================
# 设计基线（macOS HIG）：系统字体 SF Pro、系统蓝 #007AFF、浅灰分层背景、
# hairline 分隔线 #d2d2d7、6-8px 圆角控件、窄悬浮滚动条、深色 tooltip。
APP_QSS = """
* { outline: none; }
QWidget { color:#1d1d1f; }
QMainWindow { background:#ececee; }
QDialog, QMessageBox { background:#f5f5f7; }
QLabel { background:transparent; }

QPushButton {
  background:#ffffff; border:1px solid #d2d2d7; border-radius:6px;
  padding:5px 14px; color:#1d1d1f; min-height:18px;
}
QPushButton:hover { background:#f5f5f7; }
QPushButton:pressed { background:#e8e8ed; }
QPushButton:default { background:#007AFF; color:#fff; border-color:#007AFF; }
QPushButton:default:hover { background:#0071e3; }
QPushButton:disabled { color:#aeaeb2; background:#f5f5f7; border-color:#e5e5ea; }

QToolTip {
  background:rgba(45,45,48,235); color:#f5f5f7; border:none;
  border-radius:4px; padding:4px 8px; font-size:12px;
}

QMenu { background:#f5f5f7; border:1px solid #d2d2d7; border-radius:8px; padding:4px; }
QMenu::item { padding:5px 18px; border-radius:5px; }
QMenu::item:selected { background:#007AFF; color:#fff; }
QMenu::separator { height:1px; background:#d2d2d7; margin:4px 8px; }
QMenu::icon { padding-left:8px; }

QLineEdit, QPlainTextEdit {
  background:#fff; border:1px solid #d2d2d7; border-radius:6px; padding:3px 6px;
}
QLineEdit:focus, QPlainTextEdit:focus { border:1px solid #007AFF; }

QComboBox {
  background:#fff; border:1px solid #d2d2d7; border-radius:6px; padding:3px 10px 3px 8px;
}
QComboBox:hover { background:#f5f5f7; }
QComboBox::drop-down { border:none; width:22px; }
QComboBox QAbstractItemView {
  background:#f5f5f7; border:1px solid #d2d2d7; border-radius:6px;
  selection-background-color:#007AFF; selection-color:#fff; outline:none;
}

QTreeWidget, QTreeView { background:transparent; border:none; }
QTreeWidget::item, QTreeView::item { padding:4px 2px; border-radius:5px; }
QTreeWidget::item:selected, QTreeView::item:selected { background:#007AFF; color:#fff; }
QTreeWidget::item:hover, QTreeView::item:hover { background:rgba(0,0,0,0.045); }

QListWidget { background:#fff; border:1px solid #d2d2d7; border-radius:8px; }
QListWidget::item { padding:5px 8px; border-radius:5px; }
QListWidget::item:selected { background:#007AFF; color:#fff; }

QTableWidget { background:#fff; gridline-color:#e5e5ea; border:none; }
QTableWidget::item:selected { background:#d6e9ff; color:#1d1d1f; }
QHeaderView::section {
  background:#f5f5f7; color:#6e6e73; border:none; border-right:1px solid #e5e5ea;
  border-bottom:1px solid #d2d2d7; padding:4px 6px;
}

QScrollBar:vertical { background:transparent; width:9px; margin:2px; }
QScrollBar::handle:vertical { background:#c7c7cc; border-radius:3.5px; min-height:28px; }
QScrollBar::handle:vertical:hover { background:#aeaeb2; }
QScrollBar:horizontal { background:transparent; height:9px; margin:2px; }
QScrollBar::handle:horizontal { background:#c7c7cc; border-radius:3.5px; min-width:28px; }
QScrollBar::handle:horizontal:hover { background:#aeaeb2; }
QScrollBar::add-line, QScrollBar::sub-line { width:0; height:0; }
QScrollBar::add-page, QScrollBar::sub-page { background:transparent; }

QSlider::groove:horizontal { height:4px; background:#d2d2d7; border-radius:2px; }
QSlider::sub-page:horizontal { background:#007AFF; border-radius:2px; }
QSlider::handle:horizontal {
  width:16px; height:16px; margin:-6px 0; border-radius:8px;
  background:#fff; border:1px solid #d2d2d7;
}

QCheckBox::indicator {
  width:14px; height:14px; border:1px solid #d2d2d7; border-radius:4px; background:#fff;
}
QCheckBox::indicator:checked { background:#007AFF; border-color:#007AFF; }

QStatusBar { background:#f6f6f6; border-top:1px solid #d2d2d7; color:#6e6e73; }
QStatusBar QLabel { color:#6e6e73; }

QSplitter::handle { background:transparent; }
QTabWidget::pane { border:none; background:#f5f5f7; }
"""


def main():
    app = QApplication(sys.argv); app.setStyle('Fusion')
    app.setFont(FONTS.get('table'))
    app.setStyleSheet(APP_QSS)
    MainWindow().show(); sys.exit(app.exec_())

if __name__ == '__main__':
    main()
